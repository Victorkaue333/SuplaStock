from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse


def _parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value:
        try:
            return datetime.strptime(value, "%Y-%m-%d").date()
        except ValueError:
            return None
    return None


def _format_date(value) -> str:
    parsed = _parse_date(value)
    return parsed.strftime("%d/%m/%Y") if parsed else "-"


def _format_money(value) -> str:
    amount = value or Decimal("0")
    if not isinstance(amount, Decimal):
        amount = Decimal(str(amount))
    return f"R$ {amount:.2f}"


def _normalize_filename_date(value) -> str:
    parsed = _parse_date(value)
    return parsed.strftime("%Y%m%d") if parsed else "sem_data"


def _client_display(venda) -> str:
    if getattr(venda, "cliente", None):
        return venda.cliente.nome
    return getattr(venda, "cliente_nome_legado", "") or "Nao informado"


class PDFExporter:
    @staticmethod
    def exportar_vendas(vendas, data_inicio, data_fim):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            leftMargin=24,
            rightMargin=24,
            topMargin=24,
            bottomMargin=24,
        )

        total_vendas = 0
        total_lucro = Decimal("0")
        total_faturamento = Decimal("0")
        total_itens = 0

        table_data = [[
            "Data",
            "Transacao",
            "Cliente",
            "Produto",
            "Qtd",
            "Vlr. Unit.",
            "Total",
            "Lucro",
            "Status",
        ]]

        for venda in vendas:
            total_vendas += 1
            total_itens += int(venda.quantidade_vendida or 0)
            total_lucro += Decimal(str(venda.lucro or 0))
            total_faturamento += Decimal(str(venda.valor_total or 0))

            transacao = venda.transacao_id or f"LEG-{venda.id}"
            produto_nome = (venda.produto.nome or "")[:28]
            cliente_nome = _client_display(venda)[:24]

            table_data.append([
                _format_date(venda.data_venda),
                transacao[:14],
                cliente_nome,
                produto_nome,
                str(venda.quantidade_vendida),
                _format_money(venda.preco_venda_unitario),
                _format_money(venda.valor_total),
                _format_money(venda.lucro),
                venda.status_pagamento,
            ])

        styles = getSampleStyleSheet()
        story = [
            Paragraph("SuplaStock - Relatorio de Vendas", styles["Title"]),
            Paragraph(
                f"Periodo: {_format_date(data_inicio)} ate {_format_date(data_fim)}",
                styles["Normal"],
            ),
            Paragraph(
                f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
                styles["Normal"],
            ),
            Spacer(1, 10),
        ]

        table = Table(
            table_data,
            repeatRows=1,
            colWidths=[62, 80, 120, 130, 34, 74, 74, 74, 66],
        )
        table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ])
        )

        story.append(table)
        story.extend([
            Spacer(1, 10),
            Paragraph(f"Quantidade de vendas: {total_vendas}", styles["Normal"]),
            Paragraph(f"Quantidade de itens vendidos: {total_itens}", styles["Normal"]),
            Paragraph(f"Faturamento total: {_format_money(total_faturamento)}", styles["Normal"]),
            Paragraph(f"Lucro total: {_format_money(total_lucro)}", styles["Normal"]),
        ])

        doc.build(story)

        filename = (
            "relatorio_vendas_"
            f"{_normalize_filename_date(data_inicio)}_"
            f"{_normalize_filename_date(data_fim)}.pdf"
        )

        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    @staticmethod
    def exportar_estoque(produtos, tipo_exportacao='completo'):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            leftMargin=24,
            rightMargin=24,
            topMargin=24,
            bottomMargin=24,
        )

        total_produtos = 0
        total_unidades = 0
        valor_estoque = Decimal("0")

        table_data = [[
            "SKU",
            "Produto",
            "Categoria",
            "Estoque",
            "Min.",
            "Status",
            "Custo",
            "Venda",
            "Lucro",
            "Margem",
            "Validade",
        ]]

        for produto in produtos:
            total_produtos += 1
            estoque_atual = int(produto.estoque_atual or 0)
            total_unidades += estoque_atual

            preco_custo = Decimal(str(produto.preco_custo_unitario or 0))
            valor_estoque += preco_custo * Decimal(str(estoque_atual))

            estoque_minimo = int(produto.estoque_minimo or 0)
            lucro_unitario = Decimal(str(getattr(produto, 'lucro_unitario', Decimal('0')) or 0))

            if estoque_atual == 0:
                status = "Esgotado"
            elif estoque_atual <= estoque_minimo:
                status = "Critico"
            elif estoque_atual <= estoque_minimo + 3:
                status = "Baixo"
            else:
                status = "Normal"

            table_data.append([
                (getattr(produto, "sku", "") or "-")[:16],
                (produto.nome or "")[:30],
                produto.categoria.nome[:18] if produto.categoria else "Sem categoria",
                str(estoque_atual),
                str(produto.estoque_minimo),
                status,
                _format_money(produto.preco_custo_unitario),
                _format_money(produto.preco_venda_sugerido),
                _format_money(lucro_unitario),
                f"{Decimal(str(produto.margem_lucro_percentual or 0)):.2f}%",
                _format_date(produto.data_validade),
            ])

        styles = getSampleStyleSheet()
        story = [
            Paragraph("SuplaStock - Relatorio de Estoque", styles["Title"]),
            Paragraph(
                f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
                styles["Normal"],
            ),
            Spacer(1, 10),
        ]

        table = Table(table_data, repeatRows=1, colWidths=[66, 148, 82, 46, 38, 56, 60, 60, 60, 48, 66])
        table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ])
        )

        story.append(table)
        story.extend([
            Spacer(1, 10),
            Paragraph(f"Produtos ativos no relatorio: {total_produtos}", styles["Normal"]),
            Paragraph(f"Unidades em estoque: {total_unidades}", styles["Normal"]),
            Paragraph(f"Valor estimado de custo em estoque: {_format_money(valor_estoque)}", styles["Normal"]),
        ])

        doc.build(story)

        filename = f"relatorio_estoque_{tipo_exportacao}_{datetime.now().strftime('%Y%m%d')}.pdf"
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


class ExcelExporter:
    @staticmethod
    def _apply_header_style(ws, headers):
        from openpyxl.styles import Alignment, Font, PatternFill

        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        for col_index, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_index, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    @staticmethod
    def _auto_size_columns(ws, max_width=60):
        from openpyxl.utils import get_column_letter

        for idx, column in enumerate(ws.columns, start=1):
            max_length = 0
            for cell in column:
                value = cell.value
                if value is None:
                    continue
                max_length = max(max_length, len(str(value)))
            ws.column_dimensions[get_column_letter(idx)].width = min(max_length + 2, max_width)

    @staticmethod
    def exportar_vendas(vendas, data_inicio, data_fim):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Relatorio Vendas"

        headers = [
            "Data",
            "Transacao",
            "Cliente",
            "Produto",
            "Quantidade",
            "Valor Unitario",
            "Valor Total",
            "Lucro",
            "Status",
            "Forma Pagamento",
        ]
        ExcelExporter._apply_header_style(ws, headers)

        total_itens = 0
        total_faturamento = Decimal("0")
        total_lucro = Decimal("0")

        for row, venda in enumerate(vendas, start=2):
            total_itens += int(venda.quantidade_vendida or 0)
            total_faturamento += Decimal(str(venda.valor_total or 0))
            total_lucro += Decimal(str(venda.lucro or 0))

            ws.cell(row=row, column=1, value=_format_date(venda.data_venda))
            ws.cell(row=row, column=2, value=venda.transacao_id or f"LEG-{venda.id}")
            ws.cell(row=row, column=3, value=_client_display(venda))
            ws.cell(row=row, column=4, value=venda.produto.nome)
            ws.cell(row=row, column=5, value=int(venda.quantidade_vendida or 0))
            ws.cell(row=row, column=6, value=float(venda.preco_venda_unitario or 0))
            ws.cell(row=row, column=7, value=float(venda.valor_total or 0))
            ws.cell(row=row, column=8, value=float(venda.lucro or 0))
            ws.cell(row=row, column=9, value=venda.status_pagamento)
            ws.cell(row=row, column=10, value=venda.forma_pagamento_legado or "Nao informado")

            ws.cell(row=row, column=6).number_format = '"R$" #,##0.00'
            ws.cell(row=row, column=7).number_format = '"R$" #,##0.00'
            ws.cell(row=row, column=8).number_format = '"R$" #,##0.00'

        summary_row = ws.max_row + 2
        ws.cell(row=summary_row, column=1, value="RESUMO").font = Font(bold=True)
        ws.cell(row=summary_row + 1, column=1, value="Periodo").font = Font(bold=True)
        ws.cell(
            row=summary_row + 1,
            column=2,
            value=f"{_format_date(data_inicio)} ate {_format_date(data_fim)}",
        )
        ws.cell(row=summary_row + 2, column=1, value="Total de itens").font = Font(bold=True)
        ws.cell(row=summary_row + 2, column=2, value=total_itens)
        ws.cell(row=summary_row + 3, column=1, value="Faturamento total").font = Font(bold=True)
        ws.cell(row=summary_row + 3, column=2, value=float(total_faturamento))
        ws.cell(row=summary_row + 3, column=2).number_format = '"R$" #,##0.00'
        ws.cell(row=summary_row + 4, column=1, value="Lucro total").font = Font(bold=True)
        ws.cell(row=summary_row + 4, column=2, value=float(total_lucro))
        ws.cell(row=summary_row + 4, column=2).number_format = '"R$" #,##0.00'

        ExcelExporter._auto_size_columns(ws)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        filename = (
            "relatorio_vendas_"
            f"{_normalize_filename_date(data_inicio)}_"
            f"{_normalize_filename_date(data_fim)}.xlsx"
        )

        response = HttpResponse(
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    @staticmethod
    def exportar_estoque(produtos, tipo_exportacao='completo'):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Relatorio Estoque"

        headers = [
            "SKU",
            "Produto",
            "Categoria",
            "Estoque Atual",
            "Estoque Minimo",
            "Status",
            "Preco Custo",
            "Preco Venda",
            "Lucro Unitario",
            "Margem (%)",
            "Validade",
        ]
        ExcelExporter._apply_header_style(ws, headers)

        total_produtos = 0
        total_unidades = 0
        valor_estoque = Decimal("0")

        for row, produto in enumerate(produtos, start=2):
            total_produtos += 1
            estoque_atual = int(produto.estoque_atual or 0)
            estoque_minimo = int(produto.estoque_minimo or 0)
            total_unidades += estoque_atual

            preco_custo = Decimal(str(produto.preco_custo_unitario or 0))
            preco_venda = Decimal(str(produto.preco_venda_sugerido or 0))
            margem = Decimal(str(produto.margem_lucro_percentual or 0))
            valor_estoque += preco_custo * Decimal(str(estoque_atual))

            lucro_unitario = preco_venda - preco_custo

            if estoque_atual == 0:
                status = "Esgotado"
            elif estoque_atual <= estoque_minimo:
                status = "Critico"
            elif estoque_atual <= estoque_minimo + 3:
                status = "Baixo"
            else:
                status = "Normal"

            ws.cell(row=row, column=1, value=getattr(produto, 'sku', None) or '-')
            ws.cell(row=row, column=2, value=produto.nome)
            ws.cell(row=row, column=3, value=produto.categoria.nome if produto.categoria else "Sem categoria")
            ws.cell(row=row, column=4, value=estoque_atual)
            ws.cell(row=row, column=5, value=estoque_minimo)
            ws.cell(row=row, column=6, value=status)
            ws.cell(row=row, column=7, value=float(preco_custo))
            ws.cell(row=row, column=8, value=float(preco_venda))
            ws.cell(row=row, column=9, value=float(lucro_unitario))
            ws.cell(row=row, column=10, value=float(margem))
            ws.cell(row=row, column=11, value=_format_date(produto.data_validade))

            ws.cell(row=row, column=7).number_format = '"R$" #,##0.00'
            ws.cell(row=row, column=8).number_format = '"R$" #,##0.00'
            ws.cell(row=row, column=9).number_format = '"R$" #,##0.00'
            ws.cell(row=row, column=10).number_format = '0.00"%"'

        summary_row = ws.max_row + 2
        ws.cell(row=summary_row, column=1, value="RESUMO").font = Font(bold=True)
        ws.cell(row=summary_row + 1, column=1, value="Produtos no relatorio").font = Font(bold=True)
        ws.cell(row=summary_row + 1, column=2, value=total_produtos)
        ws.cell(row=summary_row + 2, column=1, value="Unidades em estoque").font = Font(bold=True)
        ws.cell(row=summary_row + 2, column=2, value=total_unidades)
        ws.cell(row=summary_row + 3, column=1, value="Valor estimado de custo").font = Font(bold=True)
        ws.cell(row=summary_row + 3, column=2, value=float(valor_estoque))
        ws.cell(row=summary_row + 3, column=2).number_format = '"R$" #,##0.00'

        ExcelExporter._auto_size_columns(ws)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"relatorio_estoque_{tipo_exportacao}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        response = HttpResponse(
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    @staticmethod
    def exportar_clientes(clientes_ranking):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Ranking Clientes"

        headers = [
            "Cliente",
            "Codigo",
            "Telefone",
            "Email",
            "Total Compras",
            "Valor Total Compras",
            "Ticket Medio",
        ]
        ExcelExporter._apply_header_style(ws, headers)

        total_clientes = 0
        total_compras = 0
        total_faturado = Decimal("0")

        for row, cliente in enumerate(clientes_ranking, start=2):
            total_clientes += 1
            compras = int(getattr(cliente, "total_compras", 0) or 0)
            valor_total = Decimal(str(getattr(cliente, "valor_total_compras", 0) or 0))
            ticket_medio = (valor_total / Decimal(str(compras))) if compras > 0 else Decimal("0")

            total_compras += compras
            total_faturado += valor_total

            ws.cell(row=row, column=1, value=cliente.nome)
            ws.cell(row=row, column=2, value=cliente.codigo or "-")
            ws.cell(row=row, column=3, value=cliente.telefone or "-")
            ws.cell(row=row, column=4, value=cliente.email or "-")
            ws.cell(row=row, column=5, value=compras)
            ws.cell(row=row, column=6, value=float(valor_total))
            ws.cell(row=row, column=7, value=float(ticket_medio))

            ws.cell(row=row, column=6).number_format = '"R$" #,##0.00'
            ws.cell(row=row, column=7).number_format = '"R$" #,##0.00'

        summary_row = ws.max_row + 2
        ws.cell(row=summary_row, column=1, value="RESUMO").font = Font(bold=True)
        ws.cell(row=summary_row + 1, column=1, value="Clientes no ranking").font = Font(bold=True)
        ws.cell(row=summary_row + 1, column=2, value=total_clientes)
        ws.cell(row=summary_row + 2, column=1, value="Total de compras").font = Font(bold=True)
        ws.cell(row=summary_row + 2, column=2, value=total_compras)
        ws.cell(row=summary_row + 3, column=1, value="Valor total comprado").font = Font(bold=True)
        ws.cell(row=summary_row + 3, column=2, value=float(total_faturado))
        ws.cell(row=summary_row + 3, column=2).number_format = '"R$" #,##0.00'

        ExcelExporter._auto_size_columns(ws)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"ranking_clientes_{datetime.now().strftime('%Y%m%d')}.xlsx"
        response = HttpResponse(
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

