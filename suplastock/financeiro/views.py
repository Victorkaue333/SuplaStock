# financeiro/views.py
from django.shortcuts import render, get_object_or_404
from django.http import HttpResponseRedirect, HttpResponse, JsonResponse
from django.urls import reverse
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db import transaction, DatabaseError, IntegrityError
from django.db.models import Q, Sum
from django.views.decorators.http import require_POST, require_GET
from django.utils import timezone
from datetime import datetime, timedelta, date
from urllib.parse import urlencode
import json
import uuid
from decimal import Decimal, InvalidOperation
from io import BytesIO
from django.core.paginator import Paginator
from usuarios.decorators import admin_requerido, vendedor_ou_admin
from estoque.models import Produto
from vendas.models import Venda, Cliente, Pagamento
from .models import ContaPagar, ContaReceber, ItemContaReceber, RecebimentoConta
from .services import (
    build_alertas_context,
    build_fluxo_caixa_context,
    build_inadimplencia_context,
    build_relatorios_context,
    parse_periodo_datas,
    registrar_movimentacao_manual as registrar_movimentacao_manual_service,
)
import logging

logger = logging.getLogger(__name__)


def _log_operational_error(event, request, exc, **context):
    user_id = getattr(getattr(request, 'user', None), 'id', None)
    logger.exception(
        'event=%s module=financeiro user_id=%s method=%s path=%s context=%s error=%s',
        event,
        user_id,
        request.method,
        request.path,
        context,
        exc.__class__.__name__,
    )


def _parse_positive_decimal(raw_value, field_name):
    try:
        value = Decimal(str(raw_value))
    except (InvalidOperation, TypeError, ValueError):
        raise ValueError(f'{field_name} invÃ¡lido.')
    if value <= 0:
        raise ValueError(f'{field_name} deve ser maior que zero.')
    return value


def _parse_positive_int(raw_value, field_name):
    try:
        value = int(raw_value)
    except (TypeError, ValueError):
        raise ValueError(f'{field_name} invÃ¡lido.')
    if value <= 0:
        raise ValueError(f'{field_name} deve ser maior que zero.')
    return value


def _is_admin_user(user):
    return (
        getattr(user, 'is_superuser', False)
        or getattr(user, 'is_admin', False)
        or getattr(user, 'is_developer', False)
    )


def _format_currency(value):
    amount = Decimal(str(value or 0)).quantize(Decimal('0.01'))
    is_negative = amount < 0
    amount = abs(amount)
    inteiro, decimal = f'{amount:.2f}'.split('.')
    inteiro_formatado = f'{int(inteiro):,}'.replace(',', '.')
    return f'{"-" if is_negative else ""}R$ {inteiro_formatado},{decimal}'


def _to_float(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _build_relatorios_kpi_payload(context):
    lucro_por_label = {
        item.get('label'): _to_float(item.get('lucro'))
        for item in context.get('lucro_por_periodo', [])
    }

    vendas_por_periodo = []
    for item in context.get('vendas_por_periodo', []):
        label = item.get('label') or '-'
        vendas_por_periodo.append({
            'label': label,
            'receita': _to_float(item.get('receita')),
            'quantidade': int(item.get('quantidade') or 0),
            'lucro': _to_float(lucro_por_label.get(label, 0)),
        })

    produtos_mais_vendidos = [
        {
            'produto': item.get('produto__nome') or '-',
            'quantidade': int(item.get('quantidade') or 0),
            'receita': _to_float(item.get('receita')),
        }
        for item in context.get('produtos_mais_vendidos', [])
    ]

    clientes_mais_compram = [
        {
            'cliente_nome': item.get('cliente__nome') or '-',
            'cliente_codigo': item.get('cliente__codigo') or '',
            'quantidade_vendas': int(item.get('quantidade_vendas') or 0),
            'total_comprado': _to_float(item.get('total_comprado')),
            'ticket_medio': _to_float(item.get('ticket_medio')),
        }
        for item in context.get('clientes_mais_compram', [])
    ]

    clientes_inadimplentes = [
        {
            'cliente_nome': item.get('cliente_nome') or '-',
            'cliente_codigo': item.get('cliente_codigo') or '',
            'valor_em_aberto': _to_float(item.get('valor_em_aberto')),
            'contas_em_atraso': int(item.get('contas_em_atraso') or 0),
            'dias_atraso': int(item.get('dias_atraso') or 0),
        }
        for item in context.get('clientes_inadimplentes', [])
    ]

    tabela_clientes = [
        {
            'cliente_nome': item.get('cliente_nome') or '-',
            'cliente_codigo': item.get('cliente_codigo') or '',
            'total_comprado': _to_float(item.get('total_comprado')),
            'total_pago': _to_float(item.get('total_pago')),
            'saldo_em_aberto': _to_float(item.get('saldo_em_aberto')),
        }
        for item in context.get('tabela_clientes', [])
    ]

    return {
        'resumo': {
            'periodo': {
                'inicio': context.get('data_inicio').strftime('%d/%m/%Y') if context.get('data_inicio') else '',
                'fim': context.get('data_fim').strftime('%d/%m/%Y') if context.get('data_fim') else '',
            },
            'total_vendas': _to_float(context.get('total_vendas')),
            'total_recebido': _to_float(context.get('total_recebido')),
            'lucro_estimado': _to_float(context.get('lucro_estimado')),
            'quantidade_vendas': int(context.get('quantidade_vendas') or 0),
            'ticket_medio': _to_float(context.get('ticket_medio')),
        },
        'vendas_por_periodo': vendas_por_periodo,
        'produtos_mais_vendidos': produtos_mais_vendidos,
        'clientes_mais_compram': clientes_mais_compram,
        'clientes_inadimplentes': clientes_inadimplentes,
        'tabela_clientes': tabela_clientes,
    }


def _build_fluxo_kpi_payload(context):
    hoje = timezone.localdate()
    mes_referencia = hoje.strftime('%Y-%m')
    movimentos_periodo = context.get('movimentacoes_periodo') or context.get('movimentacoes') or []

    def _serialize_movimento(item):
        data_mov = item.get('data')
        data_iso = data_mov.isoformat() if hasattr(data_mov, 'isoformat') else str(data_mov or '')
        data_label = data_mov.strftime('%d/%m/%Y') if hasattr(data_mov, 'strftime') else str(data_mov or '-')
        return {
            'data': data_label,
            'data_iso': data_iso,
            'tipo': item.get('tipo') or '',
            'tipo_display': item.get('tipo_display') or (item.get('tipo') or '').title(),
            'categoria': item.get('categoria') or '',
            'categoria_display': item.get('categoria_display') or item.get('categoria') or '-',
            'descricao': item.get('descricao') or '-',
            'valor': _to_float(item.get('valor')),
            'origem': item.get('origem') or '-',
            'usuario': item.get('usuario') or '-',
            'saldo_acumulado': _to_float(item.get('saldo_acumulado')),
            'sequencia': int(item.get('sequencia') or 0),
        }

    movimentos = [_serialize_movimento(item) for item in movimentos_periodo]

    def _agrupar_por_categoria(tipo):
        buckets = {}
        for item in movimentos:
            if item['tipo'] != tipo:
                continue
            categoria = item['categoria_display'] or '-'
            if categoria not in buckets:
                buckets[categoria] = {
                    'categoria': categoria,
                    'valor': 0.0,
                    'quantidade': 0,
                }
            buckets[categoria]['valor'] += _to_float(item['valor'])
            buckets[categoria]['quantidade'] += 1
        return sorted(
            buckets.values(),
            key=lambda row: (row['valor'], row['quantidade']),
            reverse=True,
        )

    serie_fluxo = []
    chart_data = context.get('chart_data') or {}
    movimento_chart = chart_data.get('movimento') or {}
    saldo_chart = chart_data.get('saldo') or {}

    labels = (
        saldo_chart.get('labels')
        or movimento_chart.get('labels')
        or chart_data.get('labels')
        or []
    )
    if not isinstance(labels, list):
        labels = []

    entradas = movimento_chart.get('entradas')
    if not isinstance(entradas, list):
        entradas = chart_data.get('entradas')
    if not isinstance(entradas, list):
        entradas = []

    saidas = movimento_chart.get('saidas')
    if not isinstance(saidas, list):
        saidas = chart_data.get('saidas')
    if not isinstance(saidas, list):
        saidas = []

    saldos = saldo_chart.get('saldos')
    if not isinstance(saldos, list):
        saldos = chart_data.get('saldo')
    if not isinstance(saldos, list):
        saldos = []
    tamanho = min(len(labels), len(entradas), len(saidas), len(saldos))
    for idx in range(tamanho):
        serie_fluxo.append({
            'label': labels[idx],
            'entrada': _to_float(entradas[idx]),
            'saida': _to_float(saidas[idx]),
            'saldo': _to_float(saldos[idx]),
        })

    movimentos_hoje = [item for item in movimentos if item['data_iso'] == hoje.isoformat()]
    movimentos_mes = [item for item in movimentos if item['data_iso'].startswith(mes_referencia)]

    entradas_hoje = sum((item['valor'] for item in movimentos_hoje if item['tipo'] == 'entrada'), 0.0)
    saidas_hoje = sum((item['valor'] for item in movimentos_hoje if item['tipo'] == 'saida'), 0.0)
    entradas_mes = sum((item['valor'] for item in movimentos_mes if item['tipo'] == 'entrada'), 0.0)
    saidas_mes = sum((item['valor'] for item in movimentos_mes if item['tipo'] == 'saida'), 0.0)

    return {
        'resumo': {
            'periodo': {
                'inicio': context.get('data_inicio').strftime('%d/%m/%Y') if context.get('data_inicio') else '',
                'fim': context.get('data_fim').strftime('%d/%m/%Y') if context.get('data_fim') else '',
            },
            'granularidade_chart': chart_data.get('granularidade_label') or 'AutomÃ¡tica',
            'hoje_iso': hoje.isoformat(),
            'mes_referencia': mes_referencia,
            'total_entradas_periodo': _to_float(context.get('total_entradas_periodo')),
            'total_saidas_periodo': _to_float(context.get('total_saidas_periodo')),
            'saldo_periodo': _to_float(context.get('saldo_periodo')),
            'saldo_inicio_periodo': _to_float(context.get('saldo_inicio_periodo')),
            'saldo_final_periodo': _to_float(context.get('saldo_final_periodo')),
            'saldo_atual': _to_float(context.get('saldo_atual')),
            'saldo_dia': _to_float(context.get('saldo_dia')),
            'saldo_mes': _to_float(context.get('saldo_mes')),
            'entradas_hoje': _to_float(entradas_hoje),
            'saidas_hoje': _to_float(saidas_hoje),
            'entradas_mes': _to_float(entradas_mes),
            'saidas_mes': _to_float(saidas_mes),
        },
        'movimentacoes': movimentos,
        'categorias_entrada': _agrupar_por_categoria('entrada'),
        'categorias_saida': _agrupar_por_categoria('saida'),
        'serie_fluxo': serie_fluxo,
    }


def _build_inadimplencia_kpi_payload(context):
    hoje = context.get('hoje') or timezone.localdate()
    clientes_contexto = context.get('clientes_inadimplentes', [])

    def _serialize_cliente(item):
        ultima_compra = item.get('ultima_compra')
        ultimo_pagamento = item.get('ultimo_pagamento')
        vencimento_antigo = item.get('vencimento_mais_antigo')
        return {
            'cliente_id': item.get('cliente_id'),
            'cliente_nome': item.get('cliente_nome') or '-',
            'cliente_codigo': item.get('cliente_codigo') or '',
            'valor_em_atraso': _to_float(item.get('valor_em_atraso')),
            'contas_em_atraso': int(item.get('contas_em_atraso') or 0),
            'dias_em_atraso': int(item.get('dias_em_atraso') or 0),
            'risco_codigo': item.get('risco_codigo') or '',
            'risco_label': item.get('risco_label') or '-',
            'ultima_compra_iso': ultima_compra.isoformat() if hasattr(ultima_compra, 'isoformat') else '',
            'ultima_compra_label': ultima_compra.strftime('%d/%m/%Y') if hasattr(ultima_compra, 'strftime') else '-',
            'ultimo_pagamento_iso': ultimo_pagamento.isoformat() if hasattr(ultimo_pagamento, 'isoformat') else '',
            'ultimo_pagamento_label': ultimo_pagamento.strftime('%d/%m/%Y') if hasattr(ultimo_pagamento, 'strftime') else '-',
            'vencimento_antigo_iso': vencimento_antigo.isoformat() if hasattr(vencimento_antigo, 'isoformat') else '',
            'vencimento_antigo_label': vencimento_antigo.strftime('%d/%m/%Y') if hasattr(vencimento_antigo, 'strftime') else '-',
        }

    clientes = [_serialize_cliente(item) for item in clientes_contexto]
    ranking_divida = sorted(
        clientes,
        key=lambda item: (item['valor_em_atraso'], item['dias_em_atraso']),
        reverse=True,
    )
    ranking_atraso = sorted(
        clientes,
        key=lambda item: (item['dias_em_atraso'], item['valor_em_atraso']),
        reverse=True,
    )

    total_contas_em_atraso = sum((item['contas_em_atraso'] for item in clientes), 0)
    total_dias = sum((item['dias_em_atraso'] for item in clientes), 0)
    media_dias_calculada = round(total_dias / len(clientes), 1) if clientes else 0
    maior_atraso = max((item['dias_em_atraso'] for item in clientes), default=0)
    menor_atraso = min((item['dias_em_atraso'] for item in clientes), default=0)

    risco_counts = {'leve': 0, 'medio': 0, 'grave': 0}
    for item in clientes:
        risco = item.get('risco_codigo')
        if risco in risco_counts:
            risco_counts[risco] += 1

    clientes_com_pagamento = sum((1 for item in clientes if item['ultimo_pagamento_iso']), 0)
    clientes_sem_pagamento = len(clientes) - clientes_com_pagamento

    maior_divida_item = ranking_divida[0] if ranking_divida else None

    return {
        'resumo': {
            'hoje_label': hoje.strftime('%d/%m/%Y') if hasattr(hoje, 'strftime') else '',
            'filtro_cliente': context.get('cliente_busca') or '',
            'total_em_atraso': _to_float(context.get('total_em_atraso')),
            'quantidade_clientes': int(context.get('quantidade_clientes_inadimplentes') or 0),
            'media_dias_atraso': _to_float(context.get('media_dias_atraso') or media_dias_calculada),
            'maior_divida': _to_float(context.get('maior_divida_em_atraso')),
            'total_contas_em_atraso': int(total_contas_em_atraso),
            'maior_atraso': int(maior_atraso),
            'menor_atraso': int(menor_atraso),
            'clientes_com_pagamento': int(clientes_com_pagamento),
            'clientes_sem_pagamento': int(clientes_sem_pagamento),
            'risco_counts': risco_counts,
            'maior_divida_cliente': maior_divida_item,
        },
        'clientes': clientes,
        'ranking_divida': ranking_divida,
        'ranking_atraso': ranking_atraso,
    }


def _build_alertas_kpi_payload(context):
    hoje = context.get('hoje') or timezone.localdate()

    def _date_label(value):
        return value.strftime('%d/%m/%Y') if hasattr(value, 'strftime') else '-'

    contas_vencendo_hoje = [
        {
            'conta_id': item.get('conta_id'),
            'cliente_id': item.get('cliente_id'),
            'cliente_nome': item.get('cliente_nome') or '-',
            'descricao': item.get('descricao') or '-',
            'valor': _to_float(item.get('valor')),
            'vencimento_iso': item.get('vencimento').isoformat() if hasattr(item.get('vencimento'), 'isoformat') else '',
            'vencimento_label': _date_label(item.get('vencimento')),
        }
        for item in context.get('contas_vencendo_hoje', [])
    ]

    contas_atrasadas = [
        {
            'conta_id': item.get('conta_id'),
            'cliente_id': item.get('cliente_id'),
            'cliente_nome': item.get('cliente_nome') or '-',
            'descricao': item.get('descricao') or '-',
            'valor': _to_float(item.get('valor')),
            'dias_atraso': int(item.get('dias_atraso') or 0),
            'vencimento_iso': item.get('vencimento').isoformat() if hasattr(item.get('vencimento'), 'isoformat') else '',
            'vencimento_label': _date_label(item.get('vencimento')),
        }
        for item in context.get('contas_atrasadas', [])
    ]

    estoque_baixo = []
    for item in context.get('estoque_baixo', []):
        falta = max(0, int(item.estoque_minimo or 0) - int(item.estoque_atual or 0))
        estoque_baixo.append({
            'produto_id': item.id,
            'produto_nome': item.nome or '-',
            'categoria': item.categoria.nome if getattr(item, 'categoria', None) else '-',
            'fornecedor': item.fornecedor.nome if getattr(item, 'fornecedor', None) else (item.fornecedor_nome_legado or '-'),
            'estoque_atual': int(item.estoque_atual or 0),
            'estoque_minimo': int(item.estoque_minimo or 0),
            'falta_para_minimo': int(falta),
            'data_validade_label': _date_label(item.data_validade),
        })

    produtos_sem_estoque = []
    for item in context.get('produtos_sem_estoque', []):
        produtos_sem_estoque.append({
            'produto_id': item.id,
            'produto_nome': item.nome or '-',
            'categoria': item.categoria.nome if getattr(item, 'categoria', None) else '-',
            'fornecedor': item.fornecedor.nome if getattr(item, 'fornecedor', None) else (item.fornecedor_nome_legado or '-'),
            'estoque_atual': int(item.estoque_atual or 0),
            'estoque_minimo': int(item.estoque_minimo or 0),
            'data_validade_label': _date_label(item.data_validade),
        })

    contas_pagar_vencendo_hoje = []
    for item in context.get('contas_pagar_vencendo_hoje', []):
        contas_pagar_vencendo_hoje.append({
            'id': item.id,
            'descricao': item.descricao or '-',
            'categoria': item.get_categoria_display() if hasattr(item, 'get_categoria_display') else '-',
            'fornecedor': item.fornecedor.nome if getattr(item, 'fornecedor', None) else '-',
            'valor': _to_float(item.valor),
            'vencimento_iso': item.data_vencimento.isoformat() if hasattr(item.data_vencimento, 'isoformat') else '',
            'vencimento_label': _date_label(item.data_vencimento),
        })

    contas_pagar_atrasadas = [
        {
            'id': item.get('id'),
            'descricao': item.get('descricao') or '-',
            'categoria': item.get('categoria') or '-',
            'valor': _to_float(item.get('valor')),
            'vencimento_iso': item.get('data_vencimento').isoformat() if hasattr(item.get('data_vencimento'), 'isoformat') else '',
            'vencimento_label': _date_label(item.get('data_vencimento')),
            'dias_atraso': int(item.get('dias_atraso') or 0),
        }
        for item in context.get('contas_pagar_atrasadas', [])
    ]

    return {
        'resumo': {
            'hoje_label': _date_label(hoje),
            'total_alertas': int(context.get('total_alertas') or 0),
            'alertas_criticos': int(context.get('alertas_criticos') or 0),
            'alertas_hoje': int(context.get('alertas_hoje') or 0),
            'total_alertas_estoque': int(context.get('total_alertas_estoque') or 0),
            'contas_vencendo_hoje_count': len(contas_vencendo_hoje),
            'contas_atrasadas_count': len(contas_atrasadas),
            'estoque_baixo_count': len(estoque_baixo),
            'produtos_sem_estoque_count': len(produtos_sem_estoque),
            'contas_pagar_vencendo_hoje_count': len(contas_pagar_vencendo_hoje),
            'contas_pagar_atrasadas_count': len(contas_pagar_atrasadas),
        },
        'contas_vencendo_hoje': contas_vencendo_hoje,
        'contas_atrasadas': contas_atrasadas,
        'estoque_baixo': estoque_baixo,
        'produtos_sem_estoque': produtos_sem_estoque,
        'contas_pagar_vencendo_hoje': contas_pagar_vencendo_hoje,
        'contas_pagar_atrasadas': contas_pagar_atrasadas,
    }


def _build_resumo_operacional_context(reference_time=None):
    now = reference_time or timezone.localtime(timezone.now())
    mes_atual = now.month
    ano_atual = now.year

    vendas_mes_qs = Venda.objects.filter(
        data_venda__month=mes_atual,
        data_venda__year=ano_atual,
    )
    vendas_mes = list(vendas_mes_qs)

    faturamento_mes = sum(
        (Decimal(str(venda.valor_total or 0)) for venda in vendas_mes),
        Decimal('0'),
    )
    lucro_mes = sum(
        (Decimal(str(venda.lucro or 0)) for venda in vendas_mes),
        Decimal('0'),
    )

    produtos_vendidos = vendas_mes_qs.aggregate(total=Sum('quantidade_vendida'))['total'] or 0
    produtos_cadastrados = Produto.objects.count()
    vendas_recentes = Venda.objects.select_related('produto', 'cliente').order_by('-data_venda')[:5]

    produto_mais_vendido = vendas_mes_qs.values('produto__nome').annotate(
        total_vendido=Sum('quantidade_vendida')
    ).order_by('-total_vendido').first()

    produtos_estoque_baixo = Produto.objects.filter(estoque_atual__lt=10).order_by('estoque_atual')
    contas_em_aberto = ContaReceber.objects.filter(status__in=['pendente', 'parcial']).count()
    contas_atrasadas = ContaReceber.objects.filter(status='atrasado').count()
    recebimentos_hoje = (
        RecebimentoConta.objects.filter(data_recebimento__date=now.date()).aggregate(total=Sum('valor'))['total']
        or Decimal('0')
    )

    return {
        'faturamento_mes': faturamento_mes,
        'lucro_mes': lucro_mes,
        'produtos_vendidos': produtos_vendidos,
        'produtos_cadastrados': produtos_cadastrados,
        'vendas_recentes': vendas_recentes,
        'produto_mais_vendido': produto_mais_vendido,
        'produtos_estoque_baixo': produtos_estoque_baixo,
        'contas_em_aberto': contas_em_aberto,
        'contas_atrasadas': contas_atrasadas,
        'recebimentos_hoje': recebimentos_hoje,
        'mes_nome': now.strftime('%B'),
        'ultima_atualizacao': now,
    }


def _build_dashboard_financeiro_context(ano_selecionado, mes_selecionado=None, *, include_chart_json=True):
    vendas_ano_qs = Venda.objects.filter(data_venda__year=ano_selecionado)
    vendas_ano = list(vendas_ano_qs)

    if mes_selecionado:
        vendas_filtradas = [venda for venda in vendas_ano if venda.data_venda.month == mes_selecionado]
        vendas_filtradas_qs = vendas_ano_qs.filter(data_venda__month=mes_selecionado)
    else:
        vendas_filtradas = vendas_ano
        vendas_filtradas_qs = vendas_ano_qs

    faturamento_total = sum(
        (Decimal(str(venda.valor_total or 0)) for venda in vendas_filtradas),
        Decimal('0'),
    )
    lucro_total = sum(
        (Decimal(str(venda.lucro or 0)) for venda in vendas_filtradas),
        Decimal('0'),
    )

    margem_lucro = (lucro_total / faturamento_total * Decimal('100')) if faturamento_total > 0 else Decimal('0')

    produto_mais_vendido = vendas_filtradas_qs.values('produto__nome').annotate(
        total=Sum('quantidade_vendida')
    ).order_by('-total').first()

    meses = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez']
    faturamento_mensal = [0.0] * 12
    lucro_mensal = [0.0] * 12

    for venda in vendas_ano:
        mes_idx = venda.data_venda.month - 1
        faturamento_mensal[mes_idx] += float(venda.valor_total or 0)
        lucro_mensal[mes_idx] += float(venda.lucro or 0)

    context = {
        'ano_selecionado': ano_selecionado,
        'mes_selecionado': mes_selecionado,
        'faturamento_total': faturamento_total,
        'lucro_total': lucro_total,
        'margem_lucro': margem_lucro,
        'produto_mais_vendido': produto_mais_vendido['produto__nome'] if produto_mais_vendido else 'N/A',
        'meses_lista': meses,
        'faturamento_mensal_lista': faturamento_mensal,
        'lucro_mensal_lista': lucro_mensal,
        'ultima_atualizacao': timezone.localtime(timezone.now()),
    }

    if include_chart_json:
        context['meses'] = json.dumps(meses)
        context['faturamento_mensal'] = json.dumps(faturamento_mensal)
        context['lucro_mensal'] = json.dumps(lucro_mensal)

    return context


def _derivar_status_por_saldo(*, saldo_aberto, total_recebido, vencimento_relevante=None, hoje=None):
    """Deriva status financeiro de forma determinÃ­stica a partir dos valores reais."""
    hoje = hoje or date.today()
    saldo = Decimal(str(saldo_aberto or 0))
    recebido = Decimal(str(total_recebido or 0))

    if saldo <= 0:
        return 'recebido'
    if vencimento_relevante and vencimento_relevante < hoje:
        return 'atrasado'
    if recebido > 0:
        return 'parcial'
    return 'pendente'


def _status_display(status):
    return dict(ContaReceber.STATUS_CHOICES).get(status, str(status).title())


def _build_cliente_conta_ledger(
    cliente,
    *,
    include_sensitive=False,
    limite_eventos=120,
    origem=None,
):
    """Monta resumo e extrato consolidado de contas a receber por cliente."""
    contas_qs = (
        ContaReceber.objects
        .filter(cliente=cliente)
        .prefetch_related('recebimentos__criado_por', 'itens__produto__categoria')
        .order_by('criado_em', 'id')
    )
    if origem:
        contas_qs = contas_qs.filter(origem=origem)
    contas_cliente = list(contas_qs)

    total_original = Decimal('0')
    total_recebido = Decimal('0')
    saldo_total_aberto = Decimal('0')
    contas_ativas = 0
    contas_atrasadas = 0
    total_lancamentos = 0
    total_pagamentos = 0
    ultima_compra = None
    primeiro_lancamento = None
    ultimo_pagamento = None
    vencimento_relevante = None
    eventos_brutos = []

    possui_fiado = False
    possui_avulsa = False

    for conta in contas_cliente:
        possui_fiado = possui_fiado or conta.origem == 'fiado'
        possui_avulsa = possui_avulsa or conta.origem == 'avulsa'

        if conta.status != 'cancelado':
            total_lancamentos += 1
            total_original += Decimal(str(conta.valor))
            recebido_conta = conta.total_recebido
            total_recebido += recebido_conta
            saldo_conta = conta.saldo_aberto
            saldo_total_aberto += saldo_conta

            if saldo_conta > 0:
                contas_ativas += 1
                if (not vencimento_relevante) or conta.data_vencimento < vencimento_relevante:
                    vencimento_relevante = conta.data_vencimento
                if conta.data_vencimento < date.today():
                    contas_atrasadas += 1

        if not ultima_compra or conta.criado_em > ultima_compra:
            ultima_compra = conta.criado_em
        if not primeiro_lancamento or conta.criado_em < primeiro_lancamento:
            primeiro_lancamento = conta.criado_em

        descricao_evento_compra = conta.descricao
        if conta.origem == 'fiado':
            detalhes_itens = []
            for item in conta.itens.all():
                nome_produto = item.produto.nome if item.produto_id else 'Produto'
                categoria_produto = (
                    item.produto.categoria.nome
                    if item.produto_id and item.produto.categoria_id
                    else 'Sem tipo'
                )
                detalhes_itens.append(f'{item.quantidade}x {nome_produto} [{categoria_produto}]')
            if detalhes_itens:
                descricao_evento_compra = f"Venda fiada: {', '.join(detalhes_itens)}"

        if conta.status != 'cancelado':
            eventos_brutos.append({
                'tipo': 'compra',
                'tipo_display': 'Compra fiado' if conta.origem == 'fiado' else 'LanÃ§amento avulso',
                'data': conta.criado_em.strftime('%d/%m/%Y %H:%M'),
                'valor': Decimal(str(conta.valor)),
                'conta_id': conta.id,
                'descricao': descricao_evento_compra,
                'status_conta': conta.status,
                'origem': conta.origem,
                'pode_excluir': conta.status not in {'recebido', 'parcial'},
                'sort_key': (conta.criado_em.date(), 0, conta.criado_em, conta.id),
            })

        recebimentos = list(conta.recebimentos.all().order_by('data_recebimento', 'criado_em', 'id'))
        for recebimento in recebimentos:
            total_pagamentos += 1
            momento_pagamento = (
                timezone.localtime(recebimento.criado_em)
                if timezone.is_aware(recebimento.criado_em)
                else recebimento.criado_em
            )
            if not ultimo_pagamento or momento_pagamento > ultimo_pagamento:
                ultimo_pagamento = momento_pagamento

            usuario_nome = ''
            if recebimento.criado_por:
                nome_completo = (
                    recebimento.criado_por.get_full_name().strip()
                    if hasattr(recebimento.criado_por, 'get_full_name')
                    else ''
                )
                usuario_nome = nome_completo or getattr(recebimento.criado_por, 'username', '') or 'UsuÃ¡rio'

            eventos_brutos.append({
                'tipo': 'pagamento',
                'tipo_display': 'Pagamento',
                'data': momento_pagamento.strftime('%d/%m/%Y %H:%M'),
                'valor': Decimal(str(recebimento.valor)),
                'conta_id': conta.id,
                'descricao': (recebimento.observacoes or '') if include_sensitive else '',
                'forma_pagamento': recebimento.forma_pagamento,
                'usuario': usuario_nome if include_sensitive else '',
                'origem': conta.origem,
                'pode_excluir': False,
                'sort_key': (
                    recebimento.data_recebimento,
                    1,
                    recebimento.criado_em,
                    recebimento.id,
                ),
            })

    eventos_brutos.sort(key=lambda item: item['sort_key'])

    saldo_corrente = Decimal('0')
    eventos = []
    for item in eventos_brutos:
        valor = Decimal(str(item['valor']))
        if item['tipo'] == 'compra':
            saldo_corrente += valor
        else:
            saldo_corrente -= valor
            if saldo_corrente < 0:
                saldo_corrente = Decimal('0')

        eventos.append({
            'tipo': item['tipo'],
            'tipo_display': item['tipo_display'],
            'data': item['data'],
            'valor': str(valor),
            'conta_id': item['conta_id'],
            'descricao': item.get('descricao', ''),
            'forma_pagamento': item.get('forma_pagamento', ''),
            'usuario': item.get('usuario', ''),
            'status_conta': item.get('status_conta', ''),
            'pode_excluir': bool(item.get('pode_excluir', False)),
            'saldo_apos': str(saldo_corrente),
        })

    if limite_eventos and limite_eventos > 0:
        eventos = eventos[-limite_eventos:]

    status = _derivar_status_por_saldo(
        saldo_aberto=saldo_total_aberto,
        total_recebido=total_recebido,
        vencimento_relevante=vencimento_relevante,
    )
    if total_lancamentos == 0 and any(c.status == 'cancelado' for c in contas_cliente):
        status = 'cancelado'

    return {
        'resumo': {
            'cliente_id': cliente.id,
            'cliente_nome': cliente.nome,
            'total_original': str(total_original),
            'total_recebido': str(total_recebido),
            'saldo_aberto': str(saldo_total_aberto if saldo_total_aberto > 0 else Decimal('0')),
            'contas_ativas': contas_ativas,
            'contas_atrasadas': contas_atrasadas,
            'total_contas': len(contas_cliente),
            'total_contas_fiado': len([c for c in contas_cliente if c.origem == 'fiado']),
            'total_lancamentos': total_lancamentos,
            'total_pagamentos': total_pagamentos,
            'ultima_compra': ultima_compra.strftime('%d/%m/%Y %H:%M') if ultima_compra else None,
            'primeiro_lancamento': primeiro_lancamento.strftime('%d/%m/%Y %H:%M') if primeiro_lancamento else None,
            'ultimo_pagamento': ultimo_pagamento.strftime('%d/%m/%Y %H:%M') if ultimo_pagamento else None,
            'vencimento_relevante': vencimento_relevante.strftime('%d/%m/%Y') if vencimento_relevante else None,
            'status': status,
            'status_display': _status_display(status),
            'possui_fiado': possui_fiado,
            'possui_avulsa': possui_avulsa,
        },
        'eventos': eventos,
        'contas_ids': [conta.id for conta in contas_cliente],
    }


def _build_cliente_fiado_ledger(cliente, *, include_sensitive=False, limite_eventos=120):
    """Compatibilidade: resumo/extrato de fiado do cliente."""
    return _build_cliente_conta_ledger(
        cliente,
        include_sensitive=include_sensitive,
        limite_eventos=limite_eventos,
        origem='fiado',
    )


@login_required
@admin_requerido
def dashboard(request):
    """View para resumo rÃ¡pido das operaÃ§Ãµes (Resumo Operacional)"""
    context = _build_resumo_operacional_context()
    return render(request, 'financeiro/resumo_operacional.html', context)


@login_required
@admin_requerido
def dashboard_financeiro(request):
    """View para dashboard financeiro"""
    try:
        ano_selecionado = int(request.GET.get('ano', timezone.localtime(timezone.now()).year))
    except (TypeError, ValueError):
        ano_selecionado = timezone.localtime(timezone.now()).year

    mes_raw = request.GET.get('mes', '')
    try:
        mes_selecionado = int(mes_raw) if mes_raw else None
    except (TypeError, ValueError):
        mes_selecionado = None
    if mes_selecionado and not 1 <= mes_selecionado <= 12:
        mes_selecionado = None

    context = _build_dashboard_financeiro_context(ano_selecionado, mes_selecionado)
    return render(request, 'financeiro/dashboard_financeiro.html', context)


@login_required
@admin_requerido
def exportar_resumo_operacional(request):
    """Exporta os dados do resumo operacional em PDF ou Excel."""
    formato = (request.GET.get('formato') or 'pdf').strip().lower()
    context = _build_resumo_operacional_context(reference_time=timezone.localtime(timezone.now()))

    if formato == 'excel':
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill

            wb = Workbook()
            ws = wb.active
            ws.title = 'Resumo Operacional'

            header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
            header_font = Font(bold=True, color='FFFFFF')

            ws['A1'] = 'SuplaStock - Resumo Operacional'
            ws['A1'].font = Font(bold=True, size=14)
            ws['A2'] = f'Ãšltima atualizaÃ§Ã£o: {context["ultima_atualizacao"].strftime("%d/%m/%Y %H:%M:%S")}'

            indicadores = [
                ('Faturamento do mÃªs', float(context['faturamento_mes'])),
                ('Lucro lÃ­quido', float(context['lucro_mes'])),
                ('Produtos vendidos', int(context['produtos_vendidos'])),
                ('Produtos cadastrados', int(context['produtos_cadastrados'])),
            ]

            ws['A4'] = 'Indicador'
            ws['B4'] = 'Valor'
            ws['A4'].fill = header_fill
            ws['B4'].fill = header_fill
            ws['A4'].font = header_font
            ws['B4'].font = header_font

            row = 5
            for nome, valor in indicadores:
                ws.cell(row=row, column=1, value=nome)
                ws.cell(row=row, column=2, value=valor)
                if nome in {'Faturamento do mÃªs', 'Lucro lÃ­quido'}:
                    ws.cell(row=row, column=2).number_format = '"R$" #,##0.00'
                row += 1

            row += 1
            ws.cell(row=row, column=1, value='Vendas recentes').font = Font(bold=True, size=12)
            row += 1
            for idx, header in enumerate(['Produto', 'Data', 'Valor'], start=1):
                cell = ws.cell(row=row, column=idx, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center')
            row += 1

            for venda in context['vendas_recentes']:
                ws.cell(row=row, column=1, value=venda.produto.nome if venda.produto_id else '-')
                ws.cell(row=row, column=2, value=venda.data_venda.strftime('%d/%m/%Y'))
                ws.cell(row=row, column=3, value=float(venda.valor_total or 0))
                ws.cell(row=row, column=3).number_format = '"R$" #,##0.00'
                row += 1

            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    value = cell.value
                    if value is None:
                        continue
                    max_length = max(max_length, len(str(value)))
                ws.column_dimensions[column].width = min(max_length + 2, 48)

            output = BytesIO()
            wb.save(output)
            output.seek(0)

            response = HttpResponse(
                output.read(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            filename = f'resumo_operacional_{context["ultima_atualizacao"].strftime("%Y%m%d_%H%M")}.xlsx'
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            return response
        except (ImportError, DatabaseError, OSError) as exc:
            _log_operational_error('exportar_resumo_operacional_excel', request, exc)
            messages.error(request, 'Erro ao exportar resumo operacional em Excel.')
            return HttpResponseRedirect(reverse('financeiro:dashboard'))

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=24,
            rightMargin=24,
            topMargin=24,
            bottomMargin=24,
        )

        styles = getSampleStyleSheet()
        story = [
            Paragraph('SuplaStock - Resumo Operacional', styles['Title']),
            Paragraph(
                f'Ãšltima atualizaÃ§Ã£o: {context["ultima_atualizacao"].strftime("%d/%m/%Y %H:%M:%S")}',
                styles['Normal'],
            ),
            Spacer(1, 10),
        ]

        indicadores = [
            ['Indicador', 'Valor'],
            ['Faturamento do mÃªs', _format_currency(context['faturamento_mes'])],
            ['Lucro lÃ­quido', _format_currency(context['lucro_mes'])],
            ['Produtos vendidos', str(context['produtos_vendidos'])],
            ['Produtos cadastrados', str(context['produtos_cadastrados'])],
        ]
        tabela_indicadores = Table(indicadores, colWidths=[250, 220], repeatRows=1)
        tabela_indicadores.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e293b')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ]))
        story.append(tabela_indicadores)
        story.append(Spacer(1, 12))

        vendas_recentes = [['Produto', 'Data', 'Valor']]
        for venda in context['vendas_recentes']:
            vendas_recentes.append([
                (venda.produto.nome if venda.produto_id else '-')[:40],
                venda.data_venda.strftime('%d/%m/%Y'),
                _format_currency(venda.valor_total),
            ])

        tabela_vendas = Table(vendas_recentes, colWidths=[250, 120, 100], repeatRows=1)
        tabela_vendas.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ]))
        story.append(Paragraph('Vendas recentes', styles['Heading3']))
        story.append(tabela_vendas)

        doc.build(story)

        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        filename = f'resumo_operacional_{context["ultima_atualizacao"].strftime("%Y%m%d_%H%M")}.pdf'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    except (ImportError, DatabaseError, OSError) as exc:
        _log_operational_error('exportar_resumo_operacional_pdf', request, exc)
        messages.error(request, 'Erro ao exportar resumo operacional em PDF.')
        return HttpResponseRedirect(reverse('financeiro:dashboard'))


@login_required
@admin_requerido
def exportar_dashboard_financeiro(request):
    """Exporta os dados do dashboard financeiro em PDF ou Excel."""
    formato = (request.GET.get('formato') or 'pdf').strip().lower()

    try:
        ano_selecionado = int(request.GET.get('ano', timezone.localtime(timezone.now()).year))
    except (TypeError, ValueError):
        ano_selecionado = timezone.localtime(timezone.now()).year

    mes_raw = request.GET.get('mes', '')
    try:
        mes_selecionado = int(mes_raw) if mes_raw else None
    except (TypeError, ValueError):
        mes_selecionado = None
    if mes_selecionado and not 1 <= mes_selecionado <= 12:
        mes_selecionado = None

    context = _build_dashboard_financeiro_context(
        ano_selecionado,
        mes_selecionado,
        include_chart_json=False,
    )

    if formato == 'excel':
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill

            wb = Workbook()
            ws_resumo = wb.active
            ws_resumo.title = 'Resumo Financeiro'
            ws_mensal = wb.create_sheet('EvoluÃ§Ã£o Mensal')

            header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
            header_font = Font(bold=True, color='FFFFFF')

            ws_resumo['A1'] = 'SuplaStock - Dashboard Financeiro'
            ws_resumo['A1'].font = Font(bold=True, size=14)
            periodo = f'{context["ano_selecionado"]}' + (
                f' / mÃªs {context["mes_selecionado"]}' if context['mes_selecionado'] else ''
            )
            ws_resumo['A2'] = f'PerÃ­odo: {periodo}'
            ws_resumo['A3'] = f'Ãšltima atualizaÃ§Ã£o: {context["ultima_atualizacao"].strftime("%d/%m/%Y %H:%M:%S")}'

            ws_resumo['A5'] = 'Indicador'
            ws_resumo['B5'] = 'Valor'
            ws_resumo['A5'].fill = header_fill
            ws_resumo['B5'].fill = header_fill
            ws_resumo['A5'].font = header_font
            ws_resumo['B5'].font = header_font

            indicadores = [
                ('Faturamento total', float(context['faturamento_total'])),
                ('Lucro total', float(context['lucro_total'])),
                ('Margem de lucro (%)', float(context['margem_lucro'])),
                ('Mais vendido', context['produto_mais_vendido']),
            ]

            row = 6
            for nome, valor in indicadores:
                ws_resumo.cell(row=row, column=1, value=nome)
                ws_resumo.cell(row=row, column=2, value=valor)
                if nome in {'Faturamento total', 'Lucro total'}:
                    ws_resumo.cell(row=row, column=2).number_format = '"R$" #,##0.00'
                if nome == 'Margem de lucro (%)':
                    ws_resumo.cell(row=row, column=2).number_format = '0.0"%"'
                row += 1

            headers = ['MÃªs', 'Faturamento', 'Lucro']
            for idx, header in enumerate(headers, start=1):
                cell = ws_mensal.cell(row=1, column=idx, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center')

            for idx, mes in enumerate(context['meses_lista'], start=2):
                ws_mensal.cell(row=idx, column=1, value=mes)
                ws_mensal.cell(row=idx, column=2, value=float(context['faturamento_mensal_lista'][idx - 2]))
                ws_mensal.cell(row=idx, column=3, value=float(context['lucro_mensal_lista'][idx - 2]))
                ws_mensal.cell(row=idx, column=2).number_format = '"R$" #,##0.00'
                ws_mensal.cell(row=idx, column=3).number_format = '"R$" #,##0.00'

            for worksheet in [ws_resumo, ws_mensal]:
                for col in worksheet.columns:
                    max_length = 0
                    column = col[0].column_letter
                    for cell in col:
                        value = cell.value
                        if value is None:
                            continue
                        max_length = max(max_length, len(str(value)))
                    worksheet.column_dimensions[column].width = min(max_length + 2, 42)

            output = BytesIO()
            wb.save(output)
            output.seek(0)

            response = HttpResponse(
                output.read(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            filename = f'dashboard_financeiro_{context["ultima_atualizacao"].strftime("%Y%m%d_%H%M")}.xlsx'
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            return response
        except (ImportError, DatabaseError, OSError) as exc:
            _log_operational_error(
                'exportar_dashboard_financeiro_excel',
                request,
                exc,
                ano=ano_selecionado,
                mes=mes_selecionado,
            )
            messages.error(request, 'Erro ao exportar dashboard financeiro em Excel.')
            return HttpResponseRedirect(reverse('financeiro:dashboard_financeiro'))

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=24,
            rightMargin=24,
            topMargin=24,
            bottomMargin=24,
        )

        styles = getSampleStyleSheet()
        periodo = f'{context["ano_selecionado"]}' + (
            f' / mÃªs {context["mes_selecionado"]}' if context['mes_selecionado'] else ''
        )
        story = [
            Paragraph('SuplaStock - Dashboard Financeiro', styles['Title']),
            Paragraph(f'PerÃ­odo: {periodo}', styles['Normal']),
            Paragraph(
                f'Ãšltima atualizaÃ§Ã£o: {context["ultima_atualizacao"].strftime("%d/%m/%Y %H:%M:%S")}',
                styles['Normal'],
            ),
            Spacer(1, 10),
        ]

        indicadores = [
            ['Indicador', 'Valor'],
            ['Faturamento total', _format_currency(context['faturamento_total'])],
            ['Lucro total', _format_currency(context['lucro_total'])],
            ['Margem de lucro', f'{Decimal(str(context["margem_lucro"])):.1f}%'],
            ['Mais vendido', context['produto_mais_vendido']],
        ]
        tabela_indicadores = Table(indicadores, colWidths=[250, 220], repeatRows=1)
        tabela_indicadores.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e293b')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ]))
        story.append(tabela_indicadores)
        story.append(Spacer(1, 12))

        mensal_data = [['MÃªs', 'Faturamento', 'Lucro']]
        for idx, mes in enumerate(context['meses_lista']):
            mensal_data.append([
                mes,
                _format_currency(context['faturamento_mensal_lista'][idx]),
                _format_currency(context['lucro_mensal_lista'][idx]),
            ])

        tabela_mensal = Table(mensal_data, colWidths=[80, 180, 180], repeatRows=1)
        tabela_mensal.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ]))
        story.append(Paragraph('EvoluÃ§Ã£o mensal', styles['Heading3']))
        story.append(tabela_mensal)

        doc.build(story)

        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        filename = f'dashboard_financeiro_{context["ultima_atualizacao"].strftime("%Y%m%d_%H%M")}.pdf'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    except (ImportError, DatabaseError, OSError) as exc:
        _log_operational_error(
            'exportar_dashboard_financeiro_pdf',
            request,
            exc,
            ano=ano_selecionado,
            mes=mes_selecionado,
        )
        messages.error(request, 'Erro ao exportar dashboard financeiro em PDF.')
        return HttpResponseRedirect(reverse('financeiro:dashboard_financeiro'))


@login_required
@admin_requerido
def relatorios_gerenciais(request):
    data_inicio, data_fim = parse_periodo_datas(
        request.GET.get('data_inicio'),
        request.GET.get('data_fim'),
    )
    tipo_relatorio = (request.GET.get('tipo') or 'geral').strip().lower()
    context = build_relatorios_context(
        data_inicio=data_inicio,
        data_fim=data_fim,
        tipo_relatorio=tipo_relatorio,
    )
    context['ultima_atualizacao'] = timezone.localtime(timezone.now())
    context['chart_data_json'] = json.dumps(context.get('chart_data') or {})
    context['kpi_details_json'] = json.dumps(_build_relatorios_kpi_payload(context))
    return render(request, 'financeiro/relatorios_gerenciais.html', context)


@login_required
@vendedor_ou_admin
def inadimplencia(request):
    cliente_busca = (request.GET.get('cliente') or '').strip()
    context = build_inadimplencia_context(cliente_busca=cliente_busca)
    context['kpi_details_json'] = json.dumps(_build_inadimplencia_kpi_payload(context))
    context['ultima_atualizacao'] = timezone.localtime(timezone.now())
    return render(request, 'financeiro/inadimplencia.html', context)


@login_required
@vendedor_ou_admin
def alertas_notificacoes(request):
    context = build_alertas_context()
    context['kpi_details_json'] = json.dumps(_build_alertas_kpi_payload(context))
    context['ultima_atualizacao'] = timezone.localtime(timezone.now())
    return render(request, 'financeiro/alertas_notificacoes.html', context)


@login_required
@admin_requerido
def fluxo_caixa(request):
    data_inicio, data_fim = parse_periodo_datas(
        request.GET.get('data_inicio'),
        request.GET.get('data_fim'),
    )
    tipo_movimentacao = (request.GET.get('tipo') or 'todos').strip().lower()
    ordenacao = (request.GET.get('ordem') or 'desc').strip().lower()

    context = build_fluxo_caixa_context(
        data_inicio=data_inicio,
        data_fim=data_fim,
        tipo_movimentacao=tipo_movimentacao,
        ordenacao=ordenacao,
    )
    context['kpi_details_json'] = json.dumps(_build_fluxo_kpi_payload(context))
    paginator = Paginator(context['movimentacoes'], 18)
    page_number = request.GET.get('page', 1)
    context['movimentacoes'] = paginator.get_page(page_number)
    context['chart_data_json'] = json.dumps(context['chart_data'])
    context['ultima_atualizacao'] = timezone.localtime(timezone.now())
    return render(request, 'financeiro/relatorios/fluxo_caixa.html', context)


@login_required
@admin_requerido
@require_POST
def registrar_movimentacao_manual(request):
    try:
        registrar_movimentacao_manual_service(
            tipo=request.POST.get('tipo_movimentacao'),
            categoria=request.POST.get('categoria'),
            descricao=request.POST.get('descricao'),
            valor_raw=request.POST.get('valor'),
            data_movimentacao_raw=request.POST.get('data_movimentacao'),
            usuario=request.user,
        )
        messages.success(request, 'MovimentaÃ§Ã£o manual registrada com sucesso.')
    except ValueError as exc:
        messages.error(request, str(exc))
    except (DatabaseError, OSError) as exc:
        _log_operational_error('registrar_movimentacao_manual', request, exc)
        messages.error(request, 'Erro ao registrar movimentaÃ§Ã£o manual.')

    params = {}
    if request.POST.get('return_data_inicio'):
        params['data_inicio'] = request.POST.get('return_data_inicio')
    if request.POST.get('return_data_fim'):
        params['data_fim'] = request.POST.get('return_data_fim')
    if request.POST.get('return_tipo'):
        params['tipo'] = request.POST.get('return_tipo')
    if request.POST.get('return_ordem'):
        params['ordem'] = request.POST.get('return_ordem')

    query_string = urlencode(params)
    redirect_url = reverse('financeiro:fluxo_caixa')
    if query_string:
        redirect_url = f'{redirect_url}?{query_string}'
    return HttpResponseRedirect(redirect_url)


# ========================= CONTAS A RECEBER =========================

@login_required
@vendedor_ou_admin
def contas_receber(request):
    """View principal para gestÃ£o de contas a receber (consolidada por cliente)."""
    hoje = date.today()
    ultima_atualizacao = timezone.localtime(timezone.now())
    ContaReceber.objects.filter(
        status__in=['pendente', 'parcial'],
        data_vencimento__lt=hoje
    ).update(status='atrasado')

    contas_base = (
        ContaReceber.objects
        .select_related('cliente', 'venda')
        .prefetch_related('itens', 'recebimentos')
        .all()
    )

    cliente_filter = (request.GET.get('cliente') or '').strip()
    if cliente_filter:
        contas_base = contas_base.filter(
            Q(cliente__nome__icontains=cliente_filter) | Q(cliente__codigo__icontains=cliente_filter)
        )

    status_filter = (request.GET.get('status') or '').strip()

    busca = (request.GET.get('busca') or '').strip()
    if busca:
        contas_base = contas_base.filter(
            Q(cliente__nome__icontains=busca) |
            Q(descricao__icontains=busca) |
            Q(cliente__codigo__icontains=busca) |
            Q(observacoes__icontains=busca)
        )

    contas_lista = list(contas_base.order_by('data_vencimento', 'id'))
    consolidado_por_cliente = {}

    for conta in contas_lista:
        cliente = conta.cliente
        linha = consolidado_por_cliente.get(cliente.id)
        if not linha:
            linha = {
                'cliente_id': cliente.id,
                'cliente_nome': cliente.nome,
                'cliente_codigo': cliente.codigo or '',
                'total_lancado': Decimal('0'),
                'total_pago': Decimal('0'),
                'saldo_aberto': Decimal('0'),
                'vencimento_relevante': None,
                'status': 'pendente',
                'status_display': _status_display('pendente'),
                'total_lancamentos': 0,
                'total_pagamentos': 0,
                'total_itens': 0,
                'possui_fiado': False,
                'possui_avulsa': False,
                'possui_nao_cancelada': False,
                'conta_referencia_id': None,
                'conta_aberta_referencia_id': None,
            }
            consolidado_por_cliente[cliente.id] = linha

        linha['possui_fiado'] = linha['possui_fiado'] or conta.origem == 'fiado'
        linha['possui_avulsa'] = linha['possui_avulsa'] or conta.origem == 'avulsa'
        linha['total_pagamentos'] += len(conta.recebimentos.all())
        linha['total_itens'] += len(conta.itens.all())

        if conta.status != 'cancelado':
            linha['possui_nao_cancelada'] = True
            linha['total_lancamentos'] += 1
            linha['total_lancado'] += Decimal(str(conta.valor))
            linha['total_pago'] += conta.total_recebido
            linha['saldo_aberto'] += conta.saldo_aberto

            if not linha['conta_referencia_id']:
                linha['conta_referencia_id'] = conta.id

            if conta.saldo_aberto > 0:
                if not linha['conta_aberta_referencia_id']:
                    linha['conta_aberta_referencia_id'] = conta.id
                if (
                    not linha['vencimento_relevante']
                    or conta.data_vencimento < linha['vencimento_relevante']
                ):
                    linha['vencimento_relevante'] = conta.data_vencimento

    linhas_todas = []
    for linha in consolidado_por_cliente.values():
        if linha['possui_nao_cancelada']:
            status_derivado = _derivar_status_por_saldo(
                saldo_aberto=linha['saldo_aberto'],
                total_recebido=linha['total_pago'],
                vencimento_relevante=linha['vencimento_relevante'],
                hoje=hoje,
            )
        else:
            status_derivado = 'cancelado'

        linha['status'] = status_derivado
        linha['status_display'] = _status_display(status_derivado)
        linha['conta_referencia_id'] = linha['conta_aberta_referencia_id'] or linha['conta_referencia_id']
        linha['descricao_resumo'] = (
            f"{linha['total_lancamentos']} lanÃ§amento(s), "
            f"{linha['total_pagamentos']} pagamento(s)"
        )
        linha['tem_divida_ativa'] = linha['saldo_aberto'] > 0 and linha['status'] != 'cancelado'
        linha['esta_atrasado'] = linha['status'] == 'atrasado'
        linhas_todas.append(linha)

    total_receber = Decimal('0')
    total_recebido = Decimal('0')
    total_pendente = 0
    total_atrasado = 0
    total_recebido_count = 0

    for linha in linhas_todas:
        if linha['status'] != 'cancelado':
            total_receber += linha['saldo_aberto']
            total_recebido += linha['total_pago']
        if linha['status'] in {'pendente', 'parcial'}:
            total_pendente += 1
        elif linha['status'] == 'atrasado':
            total_atrasado += 1
        elif linha['status'] == 'recebido' and linha['total_lancado'] > 0:
            total_recebido_count += 1

    total_contas = len(linhas_todas)
    alertas_vencimento = sum(
        1 for linha in linhas_todas
        if linha['vencimento_relevante']
        and linha['status'] in {'pendente', 'parcial'}
        and hoje <= linha['vencimento_relevante'] <= (hoje + timedelta(days=3))
    )
    vencendo_hoje = sum(
        1 for linha in linhas_todas
        if linha['vencimento_relevante'] == hoje and linha['status'] in {'pendente', 'parcial'}
    )

    vencimento_filter = (request.GET.get('vencimento') or '').strip()
    linhas_filtradas = []
    for linha in linhas_todas:
        if status_filter and linha['status'] != status_filter:
            continue

        vencimento = linha['vencimento_relevante']
        if vencimento_filter == 'proximos_7_dias':
            if not vencimento or not (hoje <= vencimento <= hoje + timedelta(days=7)):
                continue
        elif vencimento_filter == 'vencidos':
            if not vencimento or not (vencimento < hoje and linha['saldo_aberto'] > 0):
                continue
        elif vencimento_filter == 'hoje':
            if vencimento != hoje:
                continue
        elif vencimento_filter == 'mes_atual':
            if not vencimento or vencimento.month != hoje.month or vencimento.year != hoje.year:
                continue
        elif vencimento_filter == 'proximos_30_dias':
            if not vencimento or not (hoje <= vencimento <= hoje + timedelta(days=30)):
                continue

        linhas_filtradas.append(linha)

    ordem = (request.GET.get('ordem') or 'data_vencimento').strip()
    direcao = (request.GET.get('dir') or 'asc').strip().lower()
    if direcao not in {'asc', 'desc'}:
        direcao = 'asc'
    reverse_sort = direcao == 'desc'

    if ordem == 'cliente':
        linhas_filtradas.sort(key=lambda item: (item['cliente_nome'].lower(), item['cliente_id']), reverse=reverse_sort)
    elif ordem == 'valor':
        linhas_filtradas.sort(key=lambda item: (item['total_lancado'], item['cliente_nome'].lower()), reverse=reverse_sort)
    elif ordem in {'vencimento', 'data_vencimento'}:
        linhas_filtradas.sort(
            key=lambda item: (item['vencimento_relevante'] or date.max, item['cliente_nome'].lower()),
            reverse=reverse_sort,
        )
    elif ordem == 'status':
        linhas_filtradas.sort(key=lambda item: (item['status'], item['cliente_nome'].lower()), reverse=reverse_sort)
    else:
        linhas_filtradas.sort(
            key=lambda item: (item['vencimento_relevante'] or date.max, item['cliente_nome'].lower()),
            reverse=reverse_sort,
        )

    paginator = Paginator(linhas_filtradas, 10)
    page_number = request.GET.get('page', 1)
    contas_paginadas = paginator.get_page(page_number)

    query_params = request.GET.copy()
    query_params.pop('page', None)
    filtros_query_string = query_params.urlencode()

    context = {
        'contas': contas_paginadas,
        'total_receber': total_receber,
        'total_recebido': total_recebido,
        'total_pendente': total_pendente,
        'total_atrasado': total_atrasado,
        'total_recebido_count': total_recebido_count,
        'total_contas': total_contas,
        'alertas_vencimento': alertas_vencimento,
        'vencendo_hoje': vencendo_hoje,
        'cliente_filter': cliente_filter,
        'status_filter': status_filter,
        'vencimento_filter': vencimento_filter,
        'busca': busca,
        'ordem': ordem,
        'direcao': direcao,
        'clientes': Cliente.objects.filter(ativo=True).order_by('nome'),
        'vendas': Venda.objects.filter(status_pagamento__in=['Pendente', 'Parcial']).select_related('cliente', 'produto').order_by('-data_venda')[:50],
        'produtos': Produto.objects.filter(ativo=True, estoque_atual__gt=0).select_related('categoria').order_by('nome'),
        'ultima_atualizacao': ultima_atualizacao,
        'filtros_query_string': filtros_query_string,
    }

    return render(request, 'financeiro/contas_receber.html', context)


@login_required
@vendedor_ou_admin
def nova_conta_receber(request):
    """Criar nova conta a receber"""
    if request.method == 'POST':
        try:
            cliente_id = request.POST.get('cliente')
            venda_id = request.POST.get('venda', None)
            descricao = (request.POST.get('descricao') or '').strip()
            valor = _parse_positive_decimal(request.POST.get('valor'), 'Valor')
            data_vencimento = request.POST.get('data_vencimento')
            observacoes = (request.POST.get('observacoes') or '').strip()

            if not descricao:
                messages.error(request, 'DescriÃ§Ã£o Ã© obrigatÃ³ria.')
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            cliente = get_object_or_404(Cliente, id=cliente_id)
            venda = None
            if venda_id:
                venda = get_object_or_404(Venda, id=venda_id)
                if venda.cliente_id != cliente.id:
                    messages.error(request, 'A venda selecionada nÃ£o pertence ao cliente informado.')
                    return HttpResponseRedirect(reverse('financeiro:contas_receber'))
                if valor > venda.valor_pendente:
                    messages.error(request, 'Valor da conta nÃ£o pode ser maior que o saldo pendente da venda vinculada.')
                    return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            ContaReceber.objects.create(
                cliente=cliente,
                venda=venda,
                descricao=descricao,
                valor=valor,
                data_vencimento=datetime.strptime(data_vencimento, '%Y-%m-%d').date(),
                observacoes=observacoes or None,
                status='pendente',
                origem='avulsa'
            )

            messages.success(request, 'Conta a receber cadastrada com sucesso!')
            return HttpResponseRedirect(reverse('financeiro:contas_receber'))

        except ValueError as exc:
            messages.error(request, str(exc))
            return HttpResponseRedirect(reverse('financeiro:contas_receber'))
        except (DatabaseError, OSError) as exc:
            _log_operational_error(
                'nova_conta_receber',
                request,
                exc,
                cliente_id=cliente_id,
                venda_id=venda_id,
            )
            messages.error(request, 'Erro ao cadastrar conta.')
            return HttpResponseRedirect(reverse('financeiro:contas_receber'))

    clientes = Cliente.objects.filter(ativo=True).order_by('nome')
    vendas_pendentes = Venda.objects.filter(status_pagamento__in=['Pendente', 'Parcial']).select_related('cliente', 'produto')

    context = {
        'clientes': clientes,
        'vendas_pendentes': vendas_pendentes,
    }

    return render(request, 'financeiro/nova_conta_receber.html', context)


@login_required
@vendedor_ou_admin
@require_POST
def registrar_recebimento(request, conta_id):
    """Registrar recebimento de uma conta com suporte a pagamento parcial."""
    try:
        data_recebimento = request.POST.get('data_recebimento')
        forma_pagamento = request.POST.get('forma_pagamento', 'Dinheiro')
        observacoes = (request.POST.get('observacoes') or '').strip()
        valor_recebido = _parse_positive_decimal(
            request.POST.get('valor_recebido'),
            'Valor recebido',
        )
        idempotency_key = (request.POST.get('idempotency_key') or '').strip()
        data_rec = datetime.strptime(data_recebimento, '%Y-%m-%d').date()

        if not idempotency_key:
            raise ValueError('NÃ£o foi possÃ­vel validar a operaÃ§Ã£o de recebimento. Reabra o modal e tente novamente.')

        with transaction.atomic():
            conta = (
                ContaReceber.objects
                .select_for_update()
                .prefetch_related('recebimentos')
                .filter(id=conta_id)
                .first()
            )
            if not conta:
                messages.error(request, 'Conta a receber nÃ£o encontrada.')
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            if conta.status in {'recebido', 'cancelado'}:
                messages.error(request, 'A conta selecionada nÃ£o permite novo recebimento.')
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            if RecebimentoConta.objects.filter(
                conta=conta,
                idempotency_key=idempotency_key,
            ).exists():
                messages.info(
                    request,
                    'Este recebimento jÃ¡ foi processado anteriormente para esta conta.'
                )
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            # Lock dos recebimentos jÃ¡ existentes para evitar condiÃ§Ã£o de corrida.
            _ = list(conta.recebimentos.select_for_update().all())
            saldo_atual = conta.saldo_aberto
            if saldo_atual <= 0:
                conta.atualizar_status_por_recebimentos(save=True)
                messages.error(request, 'A conta jÃ¡ estÃ¡ quitada.')
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            if valor_recebido > saldo_atual:
                messages.error(
                    request,
                    f'Valor recebido ({_format_currency(valor_recebido)}) maior que saldo em aberto ({_format_currency(saldo_atual)}).',
                )
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            RecebimentoConta.objects.create(
                conta=conta,
                valor=valor_recebido,
                forma_pagamento=forma_pagamento,
                data_recebimento=data_rec,
                observacoes=observacoes or None,
                criado_por=request.user,
                idempotency_key=idempotency_key,
            )

            valor_restante = Decimal(str(valor_recebido))

            # INTEGRAÃ‡ÃƒO FIADO: distribui o recebimento entre as vendas da transaÃ§Ã£o.
            if conta.origem == 'fiado' and conta.transacao_id:
                vendas_fiadas = list(
                    Venda.objects.select_for_update().filter(transacao_id=conta.transacao_id).order_by('id')
                )
                if not vendas_fiadas:
                    raise ValueError('Conta fiada sem vendas vinculadas.')

                for venda in vendas_fiadas:
                    if valor_restante <= 0:
                        break
                    pendente = venda.valor_pendente
                    if pendente <= 0:
                        continue
                    valor_a_pagar = min(valor_restante, pendente)
                    Pagamento.objects.create(
                        cliente=conta.cliente,
                        venda=venda,
                        valor_pago=valor_a_pagar,
                        forma_pagamento=forma_pagamento,
                        data_pagamento=data_rec,
                        observacoes=f"Recebimento de conta #{conta.id}",
                    )
                    valor_restante -= valor_a_pagar

            # INTEGRAÃ‡ÃƒO AVULSA: conta vinculada a venda especÃ­fica
            elif conta.venda:
                venda = Venda.objects.select_for_update().filter(id=conta.venda_id).first()
                if venda and venda.valor_pendente > 0:
                    valor_a_pagar = min(valor_restante, venda.valor_pendente)
                    Pagamento.objects.create(
                        cliente=conta.cliente,
                        venda=venda,
                        valor_pago=valor_a_pagar,
                        forma_pagamento=forma_pagamento,
                        data_pagamento=data_rec,
                        observacoes=f"Recebimento de conta #{conta.id}",
                    )
                    valor_restante -= valor_a_pagar

            eh_avulsa_sem_venda = conta.origem == 'avulsa' and not conta.venda_id

            # Conta avulsa sem venda vinculada: lanÃ§a entrada financeira genÃ©rica.
            if valor_restante > 0:
                if eh_avulsa_sem_venda:
                    Pagamento.objects.create(
                        cliente=conta.cliente,
                        venda=None,
                        valor_pago=valor_restante,
                        forma_pagamento=forma_pagamento,
                        data_pagamento=data_rec,
                        observacoes=f"Recebimento avulso da conta #{conta.id}",
                    )
                    valor_restante = Decimal('0')
                else:
                    raise ValueError(
                        'NÃ£o foi possÃ­vel alocar integralmente o recebimento Ã s vendas vinculadas. '
                        'Verifique pendÃªncias e tente novamente.'
                    )

            conta.forma_pagamento = forma_pagamento
            if observacoes:
                timestamp = datetime.now().strftime('%d/%m/%Y %H:%M')
                conta.observacoes = (
                    (conta.observacoes or '')
                    + f'\n[{timestamp}] Recebimento: {_format_currency(valor_recebido)} ({forma_pagamento}) - {observacoes}'
                )
            conta.atualizar_status_por_recebimentos(save=True)

            saldo_final = conta.saldo_aberto
            status_label = conta.get_status_display()

        messages.success(
            request,
            f'Recebimento de {_format_currency(valor_recebido)} registrado com sucesso. '
            f'Status da conta: {status_label}. Saldo atual: {_format_currency(saldo_final)}.'
        )
        return HttpResponseRedirect(reverse('financeiro:contas_receber'))

    except ValueError as exc:
        messages.error(request, str(exc) if str(exc) else 'Dados de recebimento invÃ¡lidos.')
        return HttpResponseRedirect(reverse('financeiro:contas_receber'))
    except IntegrityError:
        messages.info(
            request,
            'Este recebimento jÃ¡ foi registrado para esta conta (operaÃ§Ã£o duplicada ignorada).'
        )
        return HttpResponseRedirect(reverse('financeiro:contas_receber'))
    except (DatabaseError, OSError) as exc:
        _log_operational_error(
            'registrar_recebimento',
            request,
            exc,
            conta_id=conta_id,
        )
        messages.error(request, 'Erro ao registrar recebimento.')
        return HttpResponseRedirect(reverse('financeiro:contas_receber'))


@login_required
@vendedor_ou_admin
@require_POST
def registrar_recebimento_cliente(request, cliente_id):
    """Registrar recebimento consolidado por cliente, rateando nas contas em aberto."""
    try:
        data_recebimento = request.POST.get('data_recebimento')
        forma_pagamento = request.POST.get('forma_pagamento', 'Dinheiro')
        observacoes = (request.POST.get('observacoes') or '').strip()
        valor_recebido = _parse_positive_decimal(
            request.POST.get('valor_recebido'),
            'Valor recebido',
        )
        idempotency_key = (request.POST.get('idempotency_key') or '').strip()
        data_rec = datetime.strptime(data_recebimento, '%Y-%m-%d').date()

        if not idempotency_key:
            raise ValueError('NÃ£o foi possÃ­vel validar a operaÃ§Ã£o de recebimento. Reabra o modal e tente novamente.')

        with transaction.atomic():
            cliente = (
                Cliente.objects
                .select_for_update()
                .filter(id=cliente_id)
                .first()
            )
            if not cliente:
                messages.error(request, 'Cliente nÃ£o encontrado para recebimento.')
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            contas_abertas = list(
                ContaReceber.objects
                .select_for_update()
                .prefetch_related('recebimentos')
                .filter(cliente_id=cliente_id)
                .exclude(status__in=['recebido', 'cancelado'])
                .order_by('data_vencimento', 'criado_em', 'id')
            )
            if not contas_abertas:
                messages.error(request, 'O cliente nÃ£o possui contas em aberto para recebimento.')
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            if RecebimentoConta.objects.filter(
                conta__cliente_id=cliente_id,
                idempotency_key=idempotency_key,
            ).exists():
                messages.info(
                    request,
                    'Este recebimento consolidado jÃ¡ foi processado anteriormente para este cliente.'
                )
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            for conta in contas_abertas:
                _ = list(conta.recebimentos.select_for_update().all())

            saldo_total = sum((conta.saldo_aberto for conta in contas_abertas), Decimal('0'))
            if saldo_total <= 0:
                for conta in contas_abertas:
                    conta.atualizar_status_por_recebimentos(save=True)
                messages.error(request, 'NÃ£o hÃ¡ saldo em aberto para o cliente selecionado.')
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            if valor_recebido > saldo_total:
                messages.error(
                    request,
                    f'Valor recebido ({_format_currency(valor_recebido)}) maior que o saldo total em aberto '
                    f'({_format_currency(saldo_total)}).',
                )
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            valor_restante = Decimal(str(valor_recebido))
            total_aplicado = Decimal('0')

            for conta in contas_abertas:
                if valor_restante <= 0:
                    break

                saldo_conta = conta.saldo_aberto
                if saldo_conta <= 0:
                    conta.atualizar_status_por_recebimentos(save=True)
                    continue

                valor_na_conta = min(valor_restante, saldo_conta)
                if valor_na_conta <= 0:
                    continue

                RecebimentoConta.objects.create(
                    conta=conta,
                    valor=valor_na_conta,
                    forma_pagamento=forma_pagamento,
                    data_recebimento=data_rec,
                    observacoes=observacoes or None,
                    criado_por=request.user,
                    idempotency_key=idempotency_key,
                )

                valor_restante_integracao = Decimal(str(valor_na_conta))

                if conta.origem == 'fiado' and conta.transacao_id:
                    vendas_fiadas = list(
                        Venda.objects.select_for_update().filter(transacao_id=conta.transacao_id).order_by('id')
                    )
                    if not vendas_fiadas:
                        raise ValueError(f'Conta fiada #{conta.id} sem vendas vinculadas.')

                    for venda in vendas_fiadas:
                        if valor_restante_integracao <= 0:
                            break
                        pendente = venda.valor_pendente
                        if pendente <= 0:
                            continue
                        valor_a_pagar = min(valor_restante_integracao, pendente)
                        Pagamento.objects.create(
                            cliente=conta.cliente,
                            venda=venda,
                            valor_pago=valor_a_pagar,
                            forma_pagamento=forma_pagamento,
                            data_pagamento=data_rec,
                            observacoes=f"Recebimento consolidado de cliente (conta #{conta.id})",
                        )
                        valor_restante_integracao -= valor_a_pagar

                elif conta.venda:
                    venda = Venda.objects.select_for_update().filter(id=conta.venda_id).first()
                    if venda and venda.valor_pendente > 0:
                        valor_a_pagar = min(valor_restante_integracao, venda.valor_pendente)
                        Pagamento.objects.create(
                            cliente=conta.cliente,
                            venda=venda,
                            valor_pago=valor_a_pagar,
                            forma_pagamento=forma_pagamento,
                            data_pagamento=data_rec,
                            observacoes=f"Recebimento consolidado de cliente (conta #{conta.id})",
                        )
                        valor_restante_integracao -= valor_a_pagar

                eh_avulsa_sem_venda = conta.origem == 'avulsa' and not conta.venda_id
                if valor_restante_integracao > 0:
                    if eh_avulsa_sem_venda:
                        Pagamento.objects.create(
                            cliente=conta.cliente,
                            venda=None,
                            valor_pago=valor_restante_integracao,
                            forma_pagamento=forma_pagamento,
                            data_pagamento=data_rec,
                            observacoes=f"Recebimento avulso consolidado (conta #{conta.id})",
                        )
                        valor_restante_integracao = Decimal('0')
                    else:
                        raise ValueError(
                            f'NÃ£o foi possÃ­vel alocar integralmente o recebimento da conta #{conta.id} '
                            'Ã s vendas vinculadas.'
                        )

                conta.forma_pagamento = forma_pagamento
                if observacoes:
                    timestamp = datetime.now().strftime('%d/%m/%Y %H:%M')
                    conta.observacoes = (
                        (conta.observacoes or '')
                        + f'\n[{timestamp}] Recebimento consolidado: {_format_currency(valor_na_conta)} '
                          f'({forma_pagamento}) - {observacoes}'
                    )
                conta.atualizar_status_por_recebimentos(save=True)

                total_aplicado += valor_na_conta
                valor_restante -= valor_na_conta

            if valor_restante > Decimal('0.0001'):
                raise ValueError(
                    'NÃ£o foi possÃ­vel distribuir totalmente o valor recebido nas contas em aberto.'
                )

            ledger = _build_cliente_conta_ledger(
                cliente,
                include_sensitive=False,
                limite_eventos=10,
            )
            saldo_final = Decimal(str(ledger['resumo'].get('saldo_aberto') or '0'))
            status_label = ledger['resumo'].get('status_display') or _status_display(
                ledger['resumo'].get('status')
            )

        messages.success(
            request,
            f'Recebimento consolidado de {_format_currency(total_aplicado)} registrado com sucesso para '
            f'{cliente.nome}. Status: {status_label}. Saldo atual: {_format_currency(saldo_final)}.'
        )
        return HttpResponseRedirect(reverse('financeiro:contas_receber'))

    except ValueError as exc:
        messages.error(request, str(exc) if str(exc) else 'Dados de recebimento invÃ¡lidos.')
        return HttpResponseRedirect(reverse('financeiro:contas_receber'))
    except IntegrityError:
        messages.info(
            request,
            'Este recebimento consolidado jÃ¡ foi registrado para este cliente (operaÃ§Ã£o duplicada ignorada).'
        )
        return HttpResponseRedirect(reverse('financeiro:contas_receber'))
    except (DatabaseError, OSError) as exc:
        _log_operational_error(
            'registrar_recebimento_cliente',
            request,
            exc,
            cliente_id=cliente_id,
        )
        messages.error(request, 'Erro ao registrar recebimento consolidado.')
        return HttpResponseRedirect(reverse('financeiro:contas_receber'))


@login_required
@vendedor_ou_admin
@require_POST
def deletar_conta_receber(request, conta_id):
    """Deletar conta a receber. Se for fiada e pendente, restaura estoque e remove vendas."""
    conta = get_object_or_404(ContaReceber, id=conta_id)
    try:
        with transaction.atomic():
            if conta.status in {'recebido', 'parcial'}:
                messages.error(
                    request,
                    'NÃ£o Ã© permitido excluir conta jÃ¡ recebida (total ou parcial).'
                )
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            # Se Ã© venda fiada e ainda pendente, restaurar estoque e remover vendas
            if conta.origem == 'fiado' and conta.transacao_id and conta.status in ['pendente', 'atrasado']:
                vendas_vinculadas = Venda.objects.select_for_update().filter(transacao_id=conta.transacao_id)
                for venda in vendas_vinculadas:
                    produto = Produto.objects.select_for_update().get(id=venda.produto_id)
                    produto.estoque_atual += venda.quantidade_vendida
                    produto.save(update_fields=['estoque_atual'])
                vendas_vinculadas.delete()

            conta.delete()
        messages.success(request, 'Conta a receber deletada com sucesso!')
    except (DatabaseError, OSError) as exc:
        _log_operational_error(
            'deletar_conta_receber',
            request,
            exc,
            conta_id=conta_id,
        )
        messages.error(request, 'Erro ao deletar conta.')

    return HttpResponseRedirect(reverse('financeiro:contas_receber'))


@login_required
@vendedor_ou_admin
def editar_conta_receber(request, conta_id):
    """Editar conta a receber existente"""
    conta = get_object_or_404(ContaReceber.objects.prefetch_related('recebimentos'), id=conta_id)

    if request.method == 'POST':
        try:
            cliente_id = request.POST.get('cliente')
            venda_id = request.POST.get('venda', None)
            descricao = (request.POST.get('descricao') or '').strip()
            valor = _parse_positive_decimal(request.POST.get('valor'), 'Valor')
            data_vencimento = request.POST.get('data_vencimento')
            observacoes = (request.POST.get('observacoes') or '').strip()
            status = request.POST.get('status', conta.status)

            if not descricao:
                messages.error(request, 'DescriÃ§Ã£o Ã© obrigatÃ³ria.')
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))
            if status not in dict(ContaReceber.STATUS_CHOICES):
                messages.error(request, 'Status invÃ¡lido.')
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            possui_recebimentos = conta.recebimentos.exists()
            if possui_recebimentos:
                # ApÃ³s qualquer recebimento, bloquear alteraÃ§Ãµes estruturais.
                if (
                    conta.cliente_id != int(cliente_id)
                    or (str(conta.venda_id or '') != str(venda_id or ''))
                    or conta.valor != valor
                ):
                    messages.error(
                        request,
                        'Conta com recebimentos nÃ£o permite alterar cliente, venda vinculada ou valor.'
                    )
                    return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            if conta.origem == 'fiado':
                # Venda fiada deve manter vÃ­nculo fechado por transaÃ§Ã£o.
                if venda_id:
                    messages.error(request, 'Conta fiada nÃ£o permite vÃ­nculo manual com venda avulsa.')
                    return HttpResponseRedirect(reverse('financeiro:contas_receber'))
                if valor != conta.valor or int(cliente_id) != conta.cliente_id:
                    messages.error(
                        request,
                        'Conta fiada nÃ£o permite alterar cliente ou valor apÃ³s criaÃ§Ã£o da transaÃ§Ã£o.'
                    )
                    return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            conta.cliente = get_object_or_404(Cliente, id=cliente_id)
            if venda_id:
                conta.venda = get_object_or_404(Venda, id=venda_id)
                if conta.venda.cliente_id != conta.cliente_id:
                    messages.error(request, 'A venda vinculada deve pertencer ao mesmo cliente da conta.')
                    return HttpResponseRedirect(reverse('financeiro:contas_receber'))
                if valor > conta.venda.valor_pendente and not possui_recebimentos:
                    messages.error(
                        request,
                        'Valor da conta nÃ£o pode ser maior que o saldo pendente da venda vinculada.'
                    )
                    return HttpResponseRedirect(reverse('financeiro:contas_receber'))
            else:
                conta.venda = None
            conta.descricao = descricao
            conta.valor = valor
            conta.data_vencimento = datetime.strptime(data_vencimento, '%Y-%m-%d').date()
            conta.observacoes = observacoes or None
            if possui_recebimentos:
                # Status apÃ³s recebimentos Ã© sempre derivado.
                conta.save(update_fields=['cliente', 'venda', 'descricao', 'valor', 'data_vencimento', 'observacoes'])
                conta.atualizar_status_por_recebimentos(save=True)
            else:
                conta.status = status
                conta.save()

            messages.success(request, 'Conta a receber atualizada com sucesso!')
            return HttpResponseRedirect(reverse('financeiro:contas_receber'))

        except ValueError as exc:
            messages.error(request, str(exc))
            return HttpResponseRedirect(reverse('financeiro:contas_receber'))
        except (DatabaseError, OSError) as exc:
            _log_operational_error(
                'editar_conta_receber',
                request,
                exc,
                conta_id=conta_id,
            )
            messages.error(request, 'Erro ao atualizar conta.')
            return HttpResponseRedirect(reverse('financeiro:contas_receber'))

    return JsonResponse({
        'id': conta.id,
        'cliente_id': conta.cliente.id,
        'venda_id': conta.venda.id if conta.venda else '',
        'descricao': conta.descricao,
        'valor': str(conta.valor),
        'data_vencimento': conta.data_vencimento.strftime('%Y-%m-%d'),
        'observacoes': conta.observacoes or '',
        'status': conta.status,
        'origem': conta.origem,
        'transacao_id': conta.transacao_id or '',
    })


# ========================= NOVA VENDA FIADA (CARRINHO) =========================

@login_required
@vendedor_ou_admin
@require_POST
def nova_venda_fiada(request):
    """Criar venda fiada com carrinho de produtos.

    Fluxo:
    1. Recebe cliente + carrinho de produtos + data de vencimento
    2. Valida estoque de todos os produtos
    3. Cria registros de Venda (status='Pendente') para cada item
    4. Deduz estoque dos produtos
    5. Cria ContaReceber vinculada via transacao_id
    6. Cria ItemContaReceber para cada produto do carrinho
    """
    try:
        payload_json = {}
        if request.content_type and 'application/json' in request.content_type:
            try:
                payload_json = json.loads((request.body or b'{}').decode('utf-8'))
            except (json.JSONDecodeError, UnicodeDecodeError):
                payload_json = {}

        cliente_id = (
            request.POST.get('cliente')
            or request.POST.get('cliente_id')
            or payload_json.get('cliente_id')
        )
        cliente_nome_informado = (
            request.POST.get('cliente_nome')
            or payload_json.get('cliente_nome')
            or ''
        ).strip()
        cliente_nome_informado = ' '.join(cliente_nome_informado.split())
        origem_conta_id_raw = (
            request.POST.get('origem_conta_id')
            or payload_json.get('origem_conta_id')
            or ''
        )
        carrinho_json = request.POST.get('carrinho_json')
        itens_json = request.POST.get('itens')
        data_vencimento = request.POST.get('data_vencimento') or payload_json.get('data_vencimento')
        observacoes = (
            (request.POST.get('observacoes') if request.method == 'POST' else None)
            or payload_json.get('observacoes')
            or ''
        ).strip()
        origem_conta_id = None

        if origem_conta_id_raw:
            origem_conta_id = _parse_positive_int(origem_conta_id_raw, 'Conta de origem')

        if not cliente_id and not cliente_nome_informado:
            messages.error(request, 'Informe o nome do cliente ou selecione um cliente cadastrado.')
            return HttpResponseRedirect(reverse('financeiro:contas_receber'))

        if not carrinho_json and not itens_json and not payload_json.get('itens'):
            messages.error(request, 'Carrinho vazio! Adicione produtos.')
            return HttpResponseRedirect(reverse('financeiro:contas_receber'))

        if carrinho_json:
            carrinho = json.loads(carrinho_json)
        elif itens_json:
            carrinho = json.loads(itens_json)
        else:
            carrinho = payload_json.get('itens', [])
        if not carrinho:
            messages.error(request, 'Carrinho vazio! Adicione produtos.')
            return HttpResponseRedirect(reverse('financeiro:contas_receber'))

        if not data_vencimento:
            raise ValueError('Informe a data de vencimento.')
        data_venc = datetime.strptime(data_vencimento, '%Y-%m-%d').date()

        carrinho_validado = []
        for item in carrinho:
            produto_id = _parse_positive_int(item.get('produto_id'), 'Produto')
            quantidade = _parse_positive_int(item.get('quantidade'), 'Quantidade')
            preco_unitario = _parse_positive_decimal(
                item.get('preco_unitario', item.get('preco')),
                'PreÃ§o unitÃ¡rio',
            )
            nome = (item.get('nome') or '').strip() or 'Produto'
            carrinho_validado.append({
                'produto_id': produto_id,
                'quantidade': quantidade,
                'preco_unitario': preco_unitario,
                'nome': nome,
            })

        # Gerar ID de transaÃ§Ã£o Ãºnico
        transacao_id = f"FIADO_{str(uuid.uuid4())[:8]}_{datetime.now().strftime('%Y%m%d%H%M%S')}"

        # Calcular valor total e montar descriÃ§Ã£o
        valor_total = Decimal('0')
        descricao_itens = []
        produtos_cache = {}
        cliente = None
        cliente_criado_automatico = False
        conta_origem = None

        with transaction.atomic():
            if cliente_id:
                cliente = Cliente.objects.filter(id=cliente_id).first()
                if not cliente:
                    raise ValueError('Cliente selecionado nÃ£o encontrado.')
            else:
                if len(cliente_nome_informado) < 2:
                    raise ValueError('Informe um nome de cliente vÃ¡lido.')
                if len(cliente_nome_informado) > 255:
                    raise ValueError('Nome do cliente deve ter no mÃ¡ximo 255 caracteres.')
                cliente = (
                    Cliente.objects
                    .filter(nome__iexact=cliente_nome_informado)
                    .order_by('id')
                    .first()
                )
                if not cliente:
                    cliente = Cliente.objects.create(nome=cliente_nome_informado)
                    cliente_criado_automatico = True

            if origem_conta_id:
                conta_origem = (
                    ContaReceber.objects
                    .select_for_update()
                    .select_related('cliente')
                    .filter(id=origem_conta_id)
                    .first()
                )
                if not conta_origem:
                    raise ValueError('Conta de origem nÃ£o encontrada para novo lanÃ§amento.')
                if conta_origem.cliente_id != cliente.id:
                    raise ValueError('A conta de origem nÃ£o pertence ao cliente informado.')

            # Validar estoque com lock pessimista
            for item in carrinho_validado:
                produto = Produto.objects.select_for_update().filter(id=item['produto_id']).first()
                if not produto:
                    raise ValueError('Produto invÃ¡lido no carrinho.')
                if produto.estoque_atual < item['quantidade']:
                    raise ValueError(
                        f'Estoque insuficiente para "{produto.nome}"! '
                        f'DisponÃ­vel: {produto.estoque_atual}, solicitado: {item["quantidade"]}.'
                    )
                produtos_cache[item['produto_id']] = produto

            for item in carrinho_validado:
                produto_ref = produtos_cache[item['produto_id']]
                subtotal = item['preco_unitario'] * item['quantidade']
                valor_total += subtotal
                categoria_produto = (
                    produto_ref.categoria.nome
                    if produto_ref.categoria_id
                    else 'Sem tipo'
                )
                descricao_itens.append(f"{item['quantidade']}x {produto_ref.nome} [{categoria_produto}]")

            descricao = f"Venda fiada: {', '.join(descricao_itens)}"
            if len(descricao) > 255:
                descricao = descricao[:252] + '...'

            observacoes_conta = observacoes if observacoes else f"Venda fiada com {len(carrinho_validado)} produto(s)"
            if conta_origem:
                observacoes_conta = (
                    f"{observacoes_conta}\n"
                    f"LanÃ§amento adicional vinculado Ã  conta #{conta_origem.id}."
                )

            conta = ContaReceber.objects.create(
                cliente=cliente,
                descricao=descricao,
                valor=valor_total,
                data_vencimento=data_venc,
                observacoes=observacoes_conta,
                status='pendente',
                origem='fiado',
                transacao_id=transacao_id
            )

            # Criar itens, vendas e deduzir estoque
            for item in carrinho_validado:
                produto = produtos_cache[item['produto_id']]
                preco_unitario = item['preco_unitario']
                quantidade = item['quantidade']

                ItemContaReceber.objects.create(
                    conta=conta,
                    produto=produto,
                    quantidade=quantidade,
                    preco_unitario=preco_unitario
                )

                Venda.objects.create(
                    produto=produto,
                    cliente=cliente,
                    quantidade_vendida=quantidade,
                    preco_venda_unitario=preco_unitario,
                    status_pagamento='Pendente',
                    transacao_id=transacao_id,
                    forma_pagamento_legado='Fiado'
                )

                produto.estoque_atual -= quantidade
                produto.save(update_fields=['estoque_atual'])

        total_itens = sum(item['quantidade'] for item in carrinho_validado)
        mensagem_contexto = (
            f'Novo lanÃ§amento adicionado ao histÃ³rico de fiado de {cliente.nome}. '
            if conta_origem else ''
        )
        messages.success(
            request,
            mensagem_contexto +
            f'Venda fiada registrada! {len(carrinho_validado)} produto(s), '
            f'{total_itens} unidade(s). Total: R$ {valor_total:.2f} â€” '
            f'Conta a receber gerada para {cliente.nome}.'
        )
        if cliente_criado_automatico:
            messages.info(request, f'Cliente "{cliente.nome}" foi cadastrado automaticamente para esta venda fiada.')
        return HttpResponseRedirect(reverse('financeiro:contas_receber'))

    except json.JSONDecodeError:
        messages.error(request, 'Erro ao processar carrinho. Tente novamente.')
        return HttpResponseRedirect(reverse('financeiro:contas_receber'))
    except ValueError as exc:
        messages.error(request, str(exc))
        return HttpResponseRedirect(reverse('financeiro:contas_receber'))
    except (DatabaseError, OSError) as exc:
        _log_operational_error(
            'nova_venda_fiada',
            request,
            exc,
            cliente_id=cliente_id,
            cliente_nome=cliente_nome_informado,
            origem_conta_id=origem_conta_id_raw,
        )
        messages.error(request, 'Erro ao registrar venda fiada.')
        return HttpResponseRedirect(reverse('financeiro:contas_receber'))


# ========================= DETALHES DA CONTA =========================

@login_required
@vendedor_ou_admin
@require_GET
def detalhe_conta_receber(request, conta_id):
    """Retorna detalhes completos de uma conta (JSON para modal de detalhes)"""
    conta = get_object_or_404(
        ContaReceber.objects.prefetch_related('itens__produto__categoria', 'recebimentos__criado_por'),
        id=conta_id
    )
    include_sensitive = _is_admin_user(request.user)

    itens = []
    for item in conta.itens.all():
        itens.append({
            'produto_nome': item.produto.nome,
            'produto_categoria': item.produto.categoria.nome if item.produto.categoria else 'Sem categoria',
            'quantidade': item.quantidade,
            'preco_unitario': str(item.preco_unitario),
            'subtotal': str(item.subtotal),
        })

    recebimentos = []
    recebimentos_qs = list(
        conta.recebimentos.all().order_by('data_recebimento', 'criado_em', 'id')
    )
    saldo_corrente = Decimal(str(conta.valor))

    for recebimento in recebimentos_qs:
        saldo_corrente -= Decimal(str(recebimento.valor))
        if saldo_corrente < 0:
            saldo_corrente = Decimal('0')

        usuario_nome = ''
        if recebimento.criado_por:
            nome_completo = (
                recebimento.criado_por.get_full_name().strip()
                if hasattr(recebimento.criado_por, 'get_full_name')
                else ''
            )
            usuario_nome = nome_completo or getattr(recebimento.criado_por, 'username', '') or 'UsuÃ¡rio'

        recebimentos.append({
            'valor': str(recebimento.valor),
            'forma_pagamento': recebimento.forma_pagamento,
            'data_recebimento': recebimento.data_recebimento.strftime('%d/%m/%Y'),
            'observacoes': (recebimento.observacoes or '') if include_sensitive else '',
            'criado_em': recebimento.criado_em.strftime('%d/%m/%Y %H:%M'),
            'usuario': usuario_nome if include_sensitive else '',
            'saldo_apos': str(saldo_corrente),
        })

    recebimentos_exibicao = list(reversed(recebimentos))
    ultimo_recebimento = recebimentos_qs[-1] if recebimentos_qs else None
    cliente_fiado_ledger = _build_cliente_conta_ledger(
        conta.cliente,
        include_sensitive=include_sensitive,
        limite_eventos=120,
    )

    payload = {
        'id': conta.id,
        'cliente_nome': conta.cliente.nome,
        'cliente_codigo': conta.cliente.codigo or '',
        'descricao': conta.descricao,
        'valor': str(conta.valor),
        'data_vencimento': conta.data_vencimento.strftime('%d/%m/%Y'),
        'data_recebimento': conta.data_recebimento.strftime('%d/%m/%Y') if conta.data_recebimento else None,
        'ultimo_pagamento': (
            timezone.localtime(ultimo_recebimento.criado_em).strftime('%d/%m/%Y %H:%M')
            if ultimo_recebimento and timezone.is_aware(ultimo_recebimento.criado_em)
            else (ultimo_recebimento.criado_em.strftime('%d/%m/%Y %H:%M') if ultimo_recebimento else None)
        ),
        'status': conta.status,
        'status_display': conta.get_status_display(),
        'origem': conta.origem,
        'origem_display': conta.get_origem_display(),
        'transacao_id': conta.transacao_id if include_sensitive else '',
        'forma_pagamento': conta.forma_pagamento or '',
        'observacoes': (conta.observacoes or '') if include_sensitive else '',
        'criado_em': conta.criado_em.strftime('%d/%m/%Y %H:%M'),
        'itens': itens,
        'recebimentos': recebimentos_exibicao,
        'total_recebido': str(conta.total_recebido),
        'saldo_aberto': str(conta.saldo_aberto),
        'is_fiado': conta.origem == 'fiado',
        'total_itens': len(itens),
        'total_pagamentos': len(recebimentos_qs),
        'cliente_fiado': cliente_fiado_ledger['resumo'],
        'extrato_cliente': cliente_fiado_ledger['eventos'],
    }
    return JsonResponse(payload)


# ========================= APIs PARA O CARRINHO =========================

@login_required
@vendedor_ou_admin
@require_GET
def api_buscar_produtos(request):
    """API JSON para buscar produtos (usado pelo carrinho via AJAX)"""
    include_sensitive = _is_admin_user(request.user)
    termo = request.GET.get('q', '').strip()
    produtos = Produto.objects.filter(ativo=True, estoque_atual__gt=0)

    if termo:
        produtos = produtos.filter(
            Q(nome__icontains=termo) |
            Q(categoria__nome__icontains=termo)
        )

    produtos = produtos.select_related('categoria').order_by('nome')[:20]

    resultado = []
    for p in produtos:
        payload = {
            'id': p.id,
            'nome': p.nome,
            'categoria': p.categoria.nome if p.categoria else 'Sem categoria',
            'preco_venda': str(p.preco_venda_sugerido),
            'estoque': p.estoque_atual,
        }
        if include_sensitive:
            payload['preco_custo'] = str(p.preco_custo_unitario)
        resultado.append(payload)

    return JsonResponse({'produtos': resultado})


@login_required
@vendedor_ou_admin
@require_GET
def api_validar_estoque(request):
    """Valida disponibilidade de estoque para uma quantidade solicitada."""
    try:
        produto_id = _parse_positive_int(request.GET.get('produto_id'), 'Produto')
        quantidade = _parse_positive_int(request.GET.get('quantidade'), 'Quantidade')
    except ValueError as exc:
        return JsonResponse({'disponivel': False, 'erro': str(exc)}, status=400)

    produto = Produto.objects.filter(id=produto_id, ativo=True).first()
    if not produto:
        return JsonResponse({'disponivel': False, 'erro': 'Produto nÃ£o encontrado.'}, status=404)

    estoque_atual = int(produto.estoque_atual or 0)
    disponivel = estoque_atual >= quantidade
    return JsonResponse({
        'disponivel': disponivel,
        'estoque_atual': estoque_atual,
        'produto_id': produto.id,
        'quantidade_solicitada': quantidade,
    })


@login_required
@vendedor_ou_admin
@require_GET
def api_buscar_clientes(request):
    """API JSON para buscar clientes (usado pelo select com busca via AJAX)"""
    include_sensitive = _is_admin_user(request.user)
    termo = request.GET.get('q', '').strip()
    clientes = Cliente.objects.filter(ativo=True)

    if termo:
        filtros = (
            Q(nome__icontains=termo) |
            Q(codigo__icontains=termo)
        )
        if include_sensitive:
            filtros |= Q(cpf__icontains=termo) | Q(telefone__icontains=termo)
        clientes = clientes.filter(filtros)

    clientes = clientes.order_by('nome')[:20]

    resultado = []
    for c in clientes:
        payload = {
            'id': c.id,
            'nome': c.nome,
            'codigo': c.codigo or '',
        }
        if include_sensitive:
            payload['telefone'] = c.telefone or ''
            payload['cpf'] = c.cpf or ''
        resultado.append(payload)

    return JsonResponse({'clientes': resultado})


@login_required
@vendedor_ou_admin
@require_GET
def api_resumo_cliente_fiado(request, cliente_id):
    """Resumo consolidado do fiado por cliente para modal de novo lanÃ§amento."""
    cliente = get_object_or_404(Cliente, id=cliente_id)
    include_sensitive = _is_admin_user(request.user)
    ledger = _build_cliente_fiado_ledger(
        cliente,
        include_sensitive=include_sensitive,
        limite_eventos=40,
    )
    return JsonResponse(ledger)


@login_required
@vendedor_ou_admin
@require_GET
def api_resumo_contas_receber(request):
    """Retorna os registros relacionados ao card de resumo clicado."""
    tipo = (request.GET.get('tipo') or '').strip().lower()
    hoje = date.today()

    contas_base = (
        ContaReceber.objects
        .select_related('cliente')
        .prefetch_related('recebimentos')
    )

    if tipo == 'pendente':
        contas_qs = contas_base.filter(
            status__in=['pendente', 'parcial'],
            data_vencimento__gte=hoje,
        ).order_by('data_vencimento', 'id')
        titulo = 'Contas Pendentes'
        subtitulo = 'Contas em aberto dentro do prazo.'
    elif tipo == 'atrasado':
        contas_qs = contas_base.exclude(
            status__in=['recebido', 'cancelado']
        ).filter(
            data_vencimento__lt=hoje
        ).order_by('data_vencimento', 'id')
        titulo = 'Contas Atrasadas'
        subtitulo = 'Contas em aberto com vencimento expirado.'
    elif tipo == 'recebido':
        contas_qs = contas_base.filter(
            status='recebido'
        ).order_by('-data_recebimento', '-data_vencimento', '-id')
        titulo = 'Contas Recebidas'
        subtitulo = 'Contas jÃ¡ quitadas.'
    elif tipo == 'total_receber':
        contas_qs = contas_base.exclude(
            status__in=['recebido', 'cancelado']
        ).order_by('data_vencimento', 'id')
        titulo = 'Total a Receber'
        subtitulo = 'Todas as contas em aberto.'
    else:
        return JsonResponse({'erro': 'Tipo de resumo invÃ¡lido.'}, status=400)

    total_itens = contas_qs.count()
    valor_total = Decimal('0')
    itens = []

    for conta in contas_qs:
        if tipo == 'recebido':
            valor_total += conta.total_recebido
        else:
            valor_total += conta.saldo_aberto

    for conta in contas_qs[:40]:
        itens.append({
            'id': conta.id,
            'cliente_nome': conta.cliente.nome,
            'cliente_codigo': conta.cliente.codigo or '',
            'descricao': conta.descricao,
            'status': conta.status,
            'status_display': conta.get_status_display(),
            'data_vencimento': conta.data_vencimento.strftime('%d/%m/%Y'),
            'valor': float(conta.valor),
            'saldo_aberto': float(conta.saldo_aberto),
            'total_recebido': float(conta.total_recebido),
            'origem': conta.origem,
            'origem_display': conta.get_origem_display(),
        })

    return JsonResponse({
        'tipo': tipo,
        'titulo': titulo,
        'subtitulo': subtitulo,
        'total_itens': total_itens,
        'valor_total': float(valor_total),
        'itens': itens,
    })


@login_required
@vendedor_ou_admin
def exportar_contas_receber(request):
    """Exportar contas a receber para PDF ou Excel"""
    import io

    formato = request.GET.get('formato', 'pdf')
    status_filter = request.GET.get('status', '')

    contas = ContaReceber.objects.select_related('cliente', 'venda').all()

    if status_filter:
        contas = contas.filter(status=status_filter)

    contas = contas.order_by('data_vencimento')

    total_geral = contas.aggregate(total=Sum('valor'))['total'] or Decimal('0')

    if formato == 'excel':
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill

            wb = Workbook()
            ws = wb.active
            ws.title = 'Contas a Receber'

            headers = ['Cliente', 'CÃ³digo', 'DescriÃ§Ã£o', 'Valor', 'Vencimento', 'Status', 'ObservaÃ§Ãµes']
            header_fill = PatternFill(start_color='667eea', end_color='667eea', fill_type='solid')
            header_font = Font(bold=True, color='FFFFFF')

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center')

            for row, conta in enumerate(contas, 2):
                ws.cell(row=row, column=1, value=conta.cliente.nome)
                ws.cell(row=row, column=2, value=conta.cliente.codigo or '-')
                ws.cell(row=row, column=3, value=conta.descricao)
                ws.cell(row=row, column=4, value=float(conta.valor))
                ws.cell(row=row, column=5, value=conta.data_vencimento.strftime('%d/%m/%Y'))
                ws.cell(row=row, column=6, value=conta.get_status_display())
                ws.cell(row=row, column=7, value=conta.observacoes or '-')

            last_row = len(contas) + 2
            ws.cell(row=last_row, column=3, value='TOTAL:').font = Font(bold=True)
            ws.cell(row=last_row, column=4, value=float(total_geral)).font = Font(bold=True)

            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except (TypeError, ValueError):
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column].width = adjusted_width

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            response = HttpResponse(
                output.read(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="contas_receber_{date.today().strftime("%Y%m%d")}.xlsx"'
            return response

        except (ImportError, DatabaseError, OSError) as exc:
            _log_operational_error(
                'exportar_contas_receber_excel',
                request,
                exc,
                status_filter=status_filter,
            )
            messages.error(request, 'Erro ao exportar Excel.')
            return HttpResponseRedirect(reverse('financeiro:contas_receber'))

    else:
        try:
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.pdfgen import canvas
            from reportlab.lib import colors
            from reportlab.platypus import Table, TableStyle

            buffer = io.BytesIO()
            p = canvas.Canvas(buffer, pagesize=landscape(A4))
            width, height = landscape(A4)

            p.setFont('Helvetica-Bold', 18)
            p.drawString(50, height - 50, 'SuplaStock - Contas a Receber')

            p.setFont('Helvetica', 10)
            p.drawString(50, height - 70, f'Gerado em: {date.today().strftime("%d/%m/%Y")}')

            data = [['Cliente', 'DescriÃ§Ã£o', 'Valor', 'Vencimento', 'Status']]

            for conta in contas[:50]:
                data.append([
                    conta.cliente.nome[:30],
                    conta.descricao[:40],
                    f'R$ {conta.valor:.2f}',
                    conta.data_vencimento.strftime('%d/%m/%Y'),
                    conta.get_status_display()
                ])

            data.append(['', '', f'TOTAL: R$ {total_geral:.2f}', '', ''])

            table = Table(data, colWidths=[150, 200, 80, 80, 80])
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#667eea')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -2), colors.white),
                ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 9),
                ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#e2e8f0')),
                ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#f8fafc')),
            ]))

            table.wrapOn(p, width, height)
            table.drawOn(p, 50, height - 120 - len(data) * 20)

            p.showPage()
            p.save()

            buffer.seek(0)
            response = HttpResponse(buffer, content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename="contas_receber_{date.today().strftime("%Y%m%d")}.pdf"'
            return response

        except (ImportError, DatabaseError, OSError) as exc:
            _log_operational_error(
                'exportar_contas_receber_pdf',
                request,
                exc,
                status_filter=status_filter,
            )
            messages.error(request, 'Erro ao exportar PDF.')
            return HttpResponseRedirect(reverse('financeiro:contas_receber'))


@login_required
@vendedor_ou_admin
@require_POST
def cancelar_conta_receber(request, conta_id):
    """Cancelar uma conta a receber. Se fiada e pendente, restaura estoque e remove vendas."""
    conta = get_object_or_404(ContaReceber, id=conta_id)
    try:
        motivo = (request.POST.get('motivo') or 'Cancelado pelo usuÃ¡rio').strip()

        with transaction.atomic():
            if conta.status == 'recebido' or conta.recebimentos.exists():
                messages.error(
                    request,
                    'NÃ£o Ã© permitido cancelar conta jÃ¡ recebida. Utilize estorno financeiro para reversÃ£o.'
                )
                return HttpResponseRedirect(reverse('financeiro:contas_receber'))

            # Se Ã© venda fiada e ainda nÃ£o foi recebida, restaurar estoque e cancelar vendas
            if conta.origem == 'fiado' and conta.transacao_id and conta.status in ['pendente', 'atrasado', 'parcial']:
                vendas_vinculadas = Venda.objects.select_for_update().filter(transacao_id=conta.transacao_id)
                for venda in vendas_vinculadas:
                    produto = Produto.objects.select_for_update().get(id=venda.produto_id)
                    produto.estoque_atual += venda.quantidade_vendida
                    produto.save(update_fields=['estoque_atual'])
                vendas_vinculadas.delete()

            conta.status = 'cancelado'
            conta.observacoes = (conta.observacoes or '') + f'\n[{datetime.now().strftime("%d/%m/%Y %H:%M")}] CANCELADO: {motivo}'
            conta.save()

        messages.warning(request, f'Conta a receber de R$ {conta.valor:.2f} cancelada.')
    except (DatabaseError, OSError) as exc:
        _log_operational_error(
            'cancelar_conta_receber',
            request,
            exc,
            conta_id=conta_id,
        )
        messages.error(request, 'Erro ao cancelar conta.')

    return HttpResponseRedirect(reverse('financeiro:contas_receber'))
